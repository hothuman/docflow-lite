"""Rebuild the synthetic sample input and two one-page Word templates."""
import csv
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "src" / "docflow" / "examples"


def base(title):
    d = Document()
    d.core_properties.author = "DocFlow Lite"
    d.core_properties.title = title
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin = s.bottom_margin = Inches(0.85)
    s.left_margin = s.right_margin = Inches(0.9)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = d.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name == "Normal" else 23 if name == "Title" else 14)
    # Remove decorative paragraph borders inherited from the default template.
    for element in d.styles.element.iter(qn("w:pBdr")):
        element.getparent().remove(element)
    d.styles["Normal"].paragraph_format.space_after = Pt(10)
    d.styles["Normal"].paragraph_format.line_spacing = 1.15
    s.header.paragraphs[0].text = "DocFlow Lite   /   Демонстрационный сценарий"
    s.header.paragraphs[0].runs[0].font.size = Pt(9)
    s.footer.paragraphs[0].text = "Все реквизиты вымышлены. Пример не является договором или бухгалтерским документом."
    s.footer.paragraphs[0].runs[0].font.size = Pt(9)
    d.add_paragraph(title, style="Title")
    return d


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    headers = ["Номер", "Клиент", "Услуга", "Сумма", "Дата", "Email"]
    rows = [
        ["001", "Демо Клиент А", "Настройка рабочих мест", "12500", "2026-10-21", "client-a@example.com"],
        ["002", "Демо «Вектор» & партнёры", "Обслуживание оборудования", "8400", "2026-10-22", "client-b@example.com"],
        ["003", "Демо Клиент В", "", "-150", "2026-10-23", "client-c@example.com"],
    ]
    with (OUT / "clients.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="163C61")
        cell.alignment = Alignment(vertical="center")
    for key, width in {"A":12,"B":36,"C":38,"D":16,"E":17,"F":32}.items():
        ws.column_dimensions[key].width = width
    ws.row_dimensions[1].height = 26
    wb.save(OUT / "clients.xlsx")
    d = base("Заявка на услугу")
    d.add_paragraph("Пример заполнения заявки из строки таблицы. Перед использованием собственного документа проверьте содержание и оформление.")
    p = d.add_paragraph("Номер заявки  ")
    # Intentional split token verifies that Word run boundaries are supported.
    p.add_run("{{invoice_").bold = True
    p.add_run("no}}")
    d.add_paragraph("Дата  {{date}}")
    d.add_heading("Данные клиента", 1)
    d.add_paragraph("Клиент  {{client}}")
    d.add_paragraph("Email  {{email}}")
    d.add_heading("Заказанная услуга", 1)
    d.add_paragraph("{{service}}")
    d.add_paragraph("Сумма по заявке  {{amount}} ₽")
    d.save(OUT / "01_request.docx")
    d = base("Подтверждение заявки")
    d.add_paragraph("Этот демонстрационный документ показывает, как второй шаблон заполняется из той же записи без повторного ввода данных.")
    d.add_paragraph("Кому  {{client}}")
    d.add_paragraph("Email  {{email}}")
    d.add_heading("Сведения о заявке", 1)
    table = d.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.1)
    table.columns[1].width = Inches(4.4)
    table.rows[0].cells[0].text = "Поле"
    table.rows[0].cells[1].text = "Значение"
    for a, b in [("Номер заявки","{{invoice_no}}"),("Дата","{{date}}"),("Услуга","{{service}}"),("Сумма","{{amount}} ₽")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = a, b
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top","left","bottom","right"):
                border = OxmlElement("w:" + edge)
                for k,v in (("val","single"),("sz","4"),("color","D9D9D9")):
                    border.set(qn("w:"+k),v)
                borders.append(border)
            tcpr.append(borders)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "163C61" if r == 0 else "F2F5F8" if r % 2 == 0 else "FFFFFF")
            tcpr.append(shading)
            margins = OxmlElement("w:tcMar")
            for edge in ("top","left","bottom","right"):
                m = OxmlElement("w:" + edge)
                m.set(qn("w:w"), "100")
                m.set(qn("w:type"), "dxa")
                margins.append(m)
            tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.space_before = Pt(5)
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255,255,255) if r == 0 else RGBColor(0,0,0)
                    run.bold = r == 0
    d.add_paragraph("")
    d.add_paragraph("Документ создан автоматически по шаблону. Перед отправкой проверьте результат.")
    d.save(OUT / "02_confirmation.docx")
    obj = {"version":1,"data_file":"clients.xlsx","templates":["01_request.docx","02_confirmation.docx"],"mapping":{"invoice_no":"Номер","client":"Клиент","service":"Услуга","amount":"Сумма","date":"Дата","email":"Email"},"rules":{"amount":{"kind":"number","required":True},"date":{"kind":"date","required":True},"email":{"kind":"email","required":True}},"name_column":"Номер"}
    (OUT / "profile.json").write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding="utf-8")
    print("Created clients.csv, clients.xlsx, 2 DOCX templates and profile.json")


if __name__ == "__main__":
    main()

