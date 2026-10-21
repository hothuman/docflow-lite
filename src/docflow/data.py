import csv
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from defusedxml import ElementTree as SafeET

from .errors import DocFlowError
from .safety import BAD_XML_CHARS, check_archive, read_bounded

MAX_ROWS = 5000
MAX_COLUMNS = 100
MAX_CELL_CHARS = 10000


def _actual_extent(raw: bytes, worksheet_path: str) -> tuple[int, int]:
    """Read physical cells before trusting producer-supplied dimension metadata.

    The path comes from openpyxl's pinned 3.1.5 ReadOnlyWorksheet. Bounds are
    checked before iter_rows, so a distant cell cannot create a huge empty grid.
    """
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        root = SafeET.fromstring(archive.read(worksheet_path), forbid_dtd=True)
    namespace = root.tag.split("}")[0] + "}"
    data = root.find(namespace + "sheetData")
    if data is None:
        raise DocFlowError("В листе нет данных")
    last_row, max_column = 0, 0
    for row in data.findall(namespace + "row"):
        row_number = int(row.get("r", last_row + 1))
        if not last_row < row_number <= MAX_ROWS + 1:
            raise DocFlowError("Лист превышает 5000 строк данных или содержит некорректные номера строк")
        last_column = 0
        for cell in row.findall(namespace + "c"):
            ref = cell.get("r")
            if ref:
                match = re.fullmatch(r"([A-Za-z]{1,3})([1-9][0-9]*)", ref)
                if not match or int(match.group(2)) != row_number:
                    raise DocFlowError("Некорректный адрес ячейки XLSX")
                column = column_index_from_string(match.group(1))
            else:
                column = last_column + 1
            if not last_column < column <= MAX_COLUMNS:
                raise DocFlowError("Лист превышает 100 столбцов или содержит повторяющиеся ячейки")
            last_column = column
            max_column = max(max_column, column)
        last_row = row_number
    return last_row, max_column


@dataclass(frozen=True)
class Record:
    row_number: int
    values: dict[str, str]


@dataclass(frozen=True)
class Table:
    source: Path
    headers: tuple[str, ...]
    records: tuple[Record, ...]


def cell_text(value: object, number_format: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    if isinstance(value, (float, int)):
        if isinstance(value, float) and not Decimal(str(value)).is_finite():
            raise DocFlowError("Неконечное числовое значение в таблице")
        if float(value).is_integer():
            text = str(int(value))
            if re.fullmatch(r"0{2,32}", number_format):
                return text.zfill(len(number_format))
            return text
        return format(Decimal(str(value)), "f")
    return str(value)


def _table(path: Path, rows) -> Table:
    try:
        _, first = next(rows)
    except StopIteration:
        raise DocFlowError("Таблица пуста") from None
    headers = tuple(x.strip() for x in first)
    if not headers or len(headers) > MAX_COLUMNS or any(not x for x in headers):
        raise DocFlowError("Первая строка должна содержать непустые заголовки; максимум 100 столбцов")
    if len(set(headers)) != len(headers):
        raise DocFlowError("Названия столбцов повторяются")
    if any(len(h) > 200 or BAD_XML_CHARS.search(h) for h in headers):
        raise DocFlowError("Некорректный заголовок столбца")
    records = []
    for row_num, cells in rows:
        if not any(x.strip() for x in cells):
            continue
        if len(cells) > len(headers):
            raise DocFlowError(f"В строке {row_num} больше значений, чем заголовков")
        if len(records) >= MAX_ROWS:
            raise DocFlowError("За один запуск поддерживается не более 5000 записей")
        cells = cells + [""] * (len(headers) - len(cells))
        for col, text in zip(headers, cells):
            if len(text) > MAX_CELL_CHARS or BAD_XML_CHARS.search(text):
                raise DocFlowError(f"Недопустимый текст или более 10000 символов: строка {row_num}, столбец {col}")
        records.append(Record(row_num, dict(zip(headers, cells))))
    if not records:
        raise DocFlowError("В таблице нет записей после заголовка")
    return Table(path.resolve(), headers, tuple(records))


def load_table(path: Path | str, *, encoding: str = "auto", sheet: str | None = None) -> Table:
    path = Path(path)
    raw = read_bounded(path)
    if path.suffix.lower() == ".csv":
        try:
            if encoding == "auto":
                try:
                    text = raw.decode("utf-8-sig")
                except UnicodeDecodeError:
                    text = raw.decode("cp1251")
            elif encoding in {"utf-8-sig", "cp1251"}:
                text = raw.decode(encoding)
            else:
                raise DocFlowError("Кодировка CSV: auto, utf-8-sig или cp1251")
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
            except csv.Error:
                delimiter = ";" if ";" in text.split("\n", 1)[0] else ","
            reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            def rows():
                for cells in reader:
                    yield reader.line_num, cells
            return _table(path, iter(rows()))
        except (UnicodeError, csv.Error) as exc:
            raise DocFlowError(f"Не удалось разобрать CSV: {exc}") from exc
    if path.suffix.lower() != ".xlsx":
        raise DocFlowError("Выберите .xlsx или .csv; .xls и .xlsm не поддерживаются")
    check_archive(raw)
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            ws = wb[sheet] if sheet else wb.active
            max_row, max_column = _actual_extent(raw, ws._worksheet_path)
            if not max_row or not max_column:
                raise DocFlowError("Лист пуст")
            ws.reset_dimensions()
            def rows():
                for n, cells in enumerate(ws.iter_rows(max_row=max_row, max_col=max_column), 1):
                    texts = []
                    for c in cells:
                        if c.data_type in {"f", "e"}:
                            raise DocFlowError(f"Формула или ошибка Excel в строке {n}; вставьте вычисленный результат как значение")
                        texts.append(cell_text(c.value, c.number_format))
                    yield n, texts
            return _table(path, iter(rows()))
        finally:
            wb.close()
    except DocFlowError:
        raise
    except Exception as exc:
        raise DocFlowError(f"Не удалось прочитать лист XLSX: {exc}") from exc

