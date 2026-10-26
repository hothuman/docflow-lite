"""Literal Service table expansion over the unchanged protected Core renderer.

Templates retain their sections, story parts, tables and embedded graphics.
Only explicit variables/row prototypes/logo markers are changed. No expressions.
"""
import bisect
from datetime import datetime
import copy
import io
import json
import re
import zipfile
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

from docx import Document
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from ..templates import Template, TOKEN, paragraph_nodes, W, parse, STORY
from ..safety import check_archive
from .attachments import safe_path, attachment_path, digest
from .finance import ITEM_KINDS
from .validation import ServiceError, rub, now, text
from .statuses import STATUSES

KINDS = {
    'intake':'Акт приёмки оборудования', 'completion':'Акт выполненных работ',
    'combined':'Карточка ремонта и выполненные работы', 'defects':'Лист дефектовки',
    'handover':'Передача оборудования', 'estimate_preliminary':'Предварительная смета',
    'estimate_final':'Окончательная смета', 'warranty':'Гарантийный талон',
}
MARKER = re.compile(r'\[\[items:(work|part|delivery|other|history)\]\]')
ITEM_FIELDS = {'item_no','item_name','item_description','item_quantity','item_unit','item_price','item_total','item_kind','item_index'}
ALIASES = {'n':'item_no','name':'item_name','qty':'item_quantity','unit':'item_unit','price':'item_price','sum':'item_total'}
ITEM_FIELDS.update(ALIASES)
HISTORY_FIELDS = {'history_number','history_date','history_client','history_complaint','history_defects',
                  'history_works','history_parts','history_total','history_status','history_notes'}
MAX_ROWS = 1000


def _row_table(document, group):
    table = document.add_table(rows=2, cols=6); table.style='Table Grid'; table.autofit=False
    widths = [.55, 2.35, .75, .65, .95, 1.15]
    for cell, label in zip(table.rows[0].cells, ['№','Наименование','Кол-во','Ед.','Цена','Сумма']):
        cell.text=label
        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'ECECF0'); cell._tc.get_or_add_tcPr().append(shade)
    repeat=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for cell, value in zip(table.rows[1].cells, ['{{n}}','[[items:'+group+']]{{name}}','{{qty}}','{{unit}}','{{price}}','{{sum}}']):
        cell.text=value
    table.rows[1].cells[1].add_paragraph('{{item_description}}')
    table.rows[1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        border=OxmlElement('w:'+side); border.set(qn('w:val'),'single'); border.set(qn('w:sz'),'4'); border.set(qn('w:color'),'D9D9D9'); borders.append(border)
    table._tbl.tblPr.append(borders)
    for column,width in zip(table.columns,widths): column.width=Inches(width)
    for row in table.rows:
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width)
            for p in cell.paragraphs: p.paragraph_format.space_after=Pt(3)
    return table


def create_templates(folder):
    """Editable neutral forms. Existing files are never overwritten."""
    folder=Path(folder); folder.mkdir(parents=True,exist_ok=True)
    for kind,title in KINDS.items():
        path=folder/(kind+'.docx')
        if path.exists(): continue
        d=Document(); section=d.sections[0]
        # Preserve the established A4 form family while extending it.
        section.page_width=Inches(8.27); section.page_height=Inches(11.69)
        section.top_margin=section.bottom_margin=Inches(.65)
        section.left_margin=section.right_margin=Inches(.65)
        for name in ('Normal','Title','Heading 1','Heading 2'):
            style=d.styles[name]; style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
            properties=style.element.find(qn('w:pPr'))
            if properties is not None:
                for border in list(properties.findall(qn('w:pBdr'))): properties.remove(border)
        d.styles['Normal'].font.size=Pt(10)
        d.styles['Normal'].paragraph_format.space_after=Pt(4)
        d.styles['Normal'].paragraph_format.line_spacing=1.0
        d.styles['Title'].font.size=Pt(20)
        d.styles['Heading 1'].font.size=Pt(12)
        d.styles['Heading 1'].paragraph_format.space_before=Pt(8)
        d.styles['Heading 1'].paragraph_format.space_after=Pt(3)
        d.add_paragraph('[[logo]]')
        d.add_paragraph(title,'Title')
        d.add_paragraph('Исполнитель: {{organization_name}}')
        d.add_paragraph('{{organization_details}}')
        d.add_paragraph('{{organization_contact}}')
        d.add_paragraph('Ремонт {{repair_number}} · Приёмка {{received_at}}')
        d.add_paragraph('Заказчик: {{client_name}}')
        d.add_paragraph('Оборудование: {{equipment_model}} · № {{equipment_number}} · серийный № {{equipment_serial}}')
        if kind in ('intake','combined'):
            d.add_paragraph('Жалоба: {{complaint}}')
            d.add_paragraph('Комплектация: {{accessories}}')
            d.add_paragraph('Внешнее состояние: {{condition}}')
        if kind in ('defects','completion','combined','estimate_preliminary','estimate_final'):
            d.add_paragraph('Дефектовка','Heading 1'); d.add_paragraph('{{defects}}')
            d.add_paragraph('План: {{plan}}')
        quote=kind.startswith('estimate_')
        if quote:
            d.add_paragraph('Версия {{estimate_version}} · Подготовлена {{estimate_date}} · {{estimate_prepared_by}}')
            d.add_paragraph('Ручная отметка состояния: {{estimate_state}}. {{estimate_comment}}')
            d.add_paragraph('Основание изменений: {{estimate_reason}}')
        if kind in ('completion','combined','estimate_preliminary','estimate_final'):
            for group,heading in (('work','Работы'),('part','Запчасти'),('delivery','Доставка'),('other','Другие согласованные расходы')):
                d.add_paragraph(('[[optional:'+group+']]' if group in ('delivery','other') else '')+heading,'Heading 1'); _row_table(d,group)
            d.add_paragraph('Работы {{work_total}} · Запчасти {{parts_total}}')
            d.add_paragraph('Доставка {{delivery_total}} · Прочие расходы {{other_total}}')
            d.add_paragraph('Скидка {{discount}} · Всего {{total}}')
            for paragraph in d.paragraphs[-3:]:paragraph.paragraph_format.keep_with_next=True
            if kind!='estimate_preliminary':
                d.add_paragraph('Внутренние записи на дату расчёта: учтено {{paid}}, остаток {{balance}}. Не банковский документ и не чек.')
            if kind=='estimate_final':
                d.add_paragraph('Предварительный согласованный расчёт: {{approved_preliminary_total}}. Разница: {{estimate_difference}}.')
        if kind=='handover':
            d.add_paragraph('Дата выдачи: {{issued_at}}')
            d.add_paragraph('Комплектация при выдаче: {{issued_accessories}}')
            d.add_paragraph('Фактические сведения о передаче: {{handover}}')
        if kind=='warranty':
            d.add_paragraph('Прототип для утверждения','Heading 1')
            d.add_paragraph('Объём и условия гарантии: {{warranty_terms}}')
            d.add_paragraph('Срок и начало действия: ____________________')
            d.add_paragraph('Окончательные условия заполняются по утверждённой форме предприятия. Программа не назначает гарантийный срок.')
        d.add_paragraph('Производственный статус: {{status}} · Примечания: {{notes}}')
        if quote:
            d.add_paragraph('Ручные отметки в программе не заменяют подтверждение клиента. Предварительная смета не является задолженностью.')
        d.add_paragraph('Проект формы. Сведения сторон и условия утверждаются отдельно; подписи не проставлены.').paragraph_format.keep_with_next=True
        d.add_paragraph('Передал ____________________     Принял ____________________')
        d.core_properties.author=''; d.core_properties.last_modified_by=''; d.core_properties.comments=''
        d.save(path)
    return {kind:folder/(kind+'.docx') for kind in KINDS}


def replace_spans(paragraph, pattern, callback):
    nodes=paragraph_nodes(paragraph); content=''.join(t.text for t in nodes); offsets=[]; offset=0
    for node in nodes: offsets.append(offset); offset+=len(node.text)
    for match in reversed(list(pattern.finditer(content))):
        first=bisect.bisect_right(offsets,match.start())-1
        last=bisect.bisect_right(offsets,match.end()-1)-1
        left=nodes[first].text[:match.start()-offsets[first]]
        right=nodes[last].text[match.end()-offsets[last]:]
        if first==last: nodes[first].text=left+callback(match)+right
        else:
            nodes[first].text=left+callback(match)
            for node in nodes[first+1:last]: node.text=''
            nodes[last].text=right


def _stories(document):
    for part in document.part.package.parts:
        if STORY.fullmatch(str(part.partname).lstrip('/')):
            yield part.element, SimpleNamespace(part=part)


def _image_width(path, section, caption=True, reserve_inches=None):
    image=Image.from_file(str(path)); width=section.page_width-section.left_margin-section.right_margin
    height=section.page_height-section.top_margin-section.bottom_margin-Inches(reserve_inches if reserve_inches is not None else (1 if caption else .25))
    if width<=0 or height<=0: raise ServiceError('В разделе шаблона нет места для изображения')
    # Integer EMU dimensions preserve aspect ratio and fit portrait/landscape.
    return int(min(width,height*image.px_width/image.px_height))


def _append_heading(document, value, level):
    # Customer templates can contain localized styles or no heading styles.
    # Apply a built-in style only when it exists; never replace their style set.
    paragraph=document.add_paragraph()
    try:
        paragraph.style=document.styles['Heading '+str(level)]
        paragraph.add_run(value)
    except KeyError:
        run=paragraph.add_run(value);run.bold=True
        run.font.size=Pt(14 if level==1 else 12)
        run.font.color.rgb=RGBColor(0,0,0)
        outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(level-1))
        paragraph._p.get_or_add_pPr().append(outline)
    paragraph.paragraph_format.keep_with_next=True
    return paragraph


def render_service(template_path, values, items, images=(), *, history=(), logo=None):
    template=Template.load(template_path)
    document=Document(io.BytesIO(template.raw)); extra={}; count=0
    stories=list(_stories(document));item_sequence=0
    for root,_ in stories:
        # Optional default expense sections disappear only when genuinely empty.
        # Custom tables without this explicit marker keep their header rows.
        for p in list(root.iter(W+'p')):
            content=''.join(t.text for t in paragraph_nodes(p))
            optional=re.search(r'\[\[optional:(delivery|other)\]\]',content)
            if optional:
                group=optional.group(1)
                if not any(item['kind']==group for item in items):
                    following=p.getnext()
                    if following is None or following.tag!=W+'tbl':
                        raise ServiceError('Необязательный раздел должен предшествовать таблице')
                    following.getparent().remove(following); p.getparent().remove(p)
                else:
                    replace_spans(p,re.compile(r'\[\[optional:(delivery|other)\]\]'),lambda m:'')
        for table in list(root.iter(W+'tbl')):
            for row in list(table.findall(W+'tr')):
                paragraphs=[p for p in row.iter(W+'p') if next(p.iterancestors(W+'tr'),None) is row]
                content=''.join(''.join(t.text for t in paragraph_nodes(p)) for p in paragraphs)
                matches=list(MARKER.finditer(content))
                if not matches: continue
                if len(matches)!=1 or row.find('.//'+W+'vMerge') is not None or row.find('.//'+W+'tbl') is not None:
                    raise ServiceError('Прототип повторяемой строки должен быть без вертикального объединения и вложенной таблицы')
                group=matches[0].group(1)
                subset=list(history) if group=='history' else [x for x in items if x['kind']==group]
                if len(subset)>MAX_ROWS: raise ServiceError('В одном DOCX не более 1000 повторяемых строк каждого типа')
                for index,item in enumerate(subset,1):
                    count+=1; clone=copy.deepcopy(row)
                    if group=='history': local=item; fields=HISTORY_FIELDS
                    else:
                        item_sequence+=1
                        local={'item_no':str(index),'item_index':str(item_sequence),'item_name':item['name'],'item_description':item['description'],
                               'item_quantity':item['quantity'],'item_unit':item['unit'],'item_price':rub(item['price_cents'] if item.get('price_known',True) else None),
                               'item_total':rub(item['total_cents'] if item.get('price_known',True) else None),'item_kind':ITEM_KINDS[item['kind']]}
                        fields=ITEM_FIELDS
                    for p in clone.iter(W+'p'):
                        replace_spans(p,MARKER,lambda m:'')
                        def rename(match):
                            field=match.group(1)
                            if field not in fields: return match.group(0)
                            name=f'_s{count}_{field}'; extra[name]=str(local[ALIASES.get(field,field)])
                            return '{{'+name+'}}'
                        replace_spans(p,TOKEN,rename)
                    row.addprevious(clone)
                table.remove(row)
        for p in root.iter(W+'p'):
            if '[[items:' in ''.join(t.text for t in paragraph_nodes(p)):
                raise ServiceError('Неподдерживаемый или неправильно размещённый маркер строк')
    for root,owner in stories:
        for p in list(root.iter(W+'p')):
            if '[[logo]]' in ''.join(t.text for t in paragraph_nodes(p)):
                replace_spans(p,re.compile(r'\[\[logo\]\]'),lambda m:'')
                if logo:
                    paragraph=Paragraph(p,owner)
                    paragraph.add_run().add_picture(str(logo),width=Inches(1.25))
                elif not ''.join(t.text for t in paragraph_nodes(p)):
                    p.getparent().remove(p)
    if len(images)>10: raise ServiceError('Для документа выберите не более десяти изображений')
    if images:
        _append_heading(document,'Приложение с изображениями',1).paragraph_format.page_break_before=True
        previous_stage=None
        for index,image in enumerate(images):
            path,caption=image[:2]; stage=image[2] if len(image)>2 else ''
            if Path(path).stat().st_size>10*1024**2:
                raise ServiceError('Изображение для DOCX больше 10 MiB; подготовьте уменьшенную копию')
            if stage and stage!=previous_stage:
                heading=_append_heading(document,stage,2)
                if index:heading.paragraph_format.page_break_before=True
                previous_stage=stage
            else:heading=None
            name=f'_image_caption_{index}'; extra[name]=text(caption,'Подпись изображения',limit=2000)
            paragraph=document.add_paragraph('{{'+name+'}}'); paragraph.paragraph_format.keep_with_next=True
            if index and heading is None:paragraph.paragraph_format.page_break_before=True
            picture=document.add_paragraph(); picture.paragraph_format.keep_together=True
            picture.paragraph_format.line_spacing=1
            picture.paragraph_format.space_before=picture.paragraph_format.space_after=Pt(0)
            section=document.sections[-1]
            width_inches=(section.page_width-section.left_margin-section.right_margin)/Inches(1)
            # Leave space for the heading, stage and wrapped caption. Exact
            # Word pagination depends on the template's fonts; visual QA still
            # remains required for a customer's final form.
            caption_lines=1+len(extra[name])//max(20,int(width_inches*11))
            reserve=1.1+(0.5 if index==0 else 0)+caption_lines*.2
            picture.add_run().add_picture(str(path),width=_image_width(path,section,reserve_inches=reserve))
    raw=io.BytesIO(); document.save(raw); raw=raw.getvalue(); check_archive(raw)
    fields=set()
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        for name in z.namelist():
            if STORY.fullmatch(name):
                root=parse(z.read(name))
                for p in root.iter(W+'p'):
                    fields.update(m.group(1) for m in TOKEN.finditer(''.join(t.text for t in paragraph_nodes(p))))
    available={**values,**extra}
    if fields-available.keys(): raise ServiceError('Неизвестные поля шаблона: '+', '.join(sorted(fields-available.keys())))
    return Template(Path(template_path),raw,frozenset(fields)).render(available)


def _history(store, repair):
    rows=store.repairs(equipment_id=repair['equipment_id'],ascending=True)
    with store.db.connection() as c:
        items={}
        for row in c.execute('SELECT i.* FROM repair_items i JOIN repair_orders r ON i.repair_id=r.id WHERE r.equipment_id=? ORDER BY i.position,i.rowid',(repair['equipment_id'],)):
            items.setdefault(row['repair_id'],[]).append(dict(row))
    result=[]
    for row in rows:
        if row['id']==repair['id']: continue
        def listing(kind): return '\n'.join(x['name']+(': '+x['description'] if x['description'] and x['description']!=x['name'] else '') for x in items.get(row['id'],[]) if x['kind']==kind)
        result.append({'history_number':row['number'],'history_date':row['received_at'],'history_client':row['client'],
                       'history_complaint':row['complaint'],'history_defects':row['defects'],'history_works':listing('work'),
                       'history_parts':listing('part'),'history_total':rub(row['total']),'history_status':row['status_display'],'history_notes':row['notes']})
    return result


def generate_document(store, repair_id, kind, template=None, image_ids=(), *, estimate_id=None, image_choices=None):
    if kind not in KINDS: raise ServiceError('Неизвестный документ')
    with store.lock:
        snapshot=store.snapshot(repair_id); r=snapshot['repair']; eq=snapshot['equipment']; client=snapshot['client']
        amounts=snapshot['totals']; organization=snapshot['profile']; items=snapshot['repair_items']; quote=None
        if kind.startswith('estimate_'):
            from .estimates import get_estimate
            expected='preliminary' if kind=='estimate_preliminary' else 'final'
            if not estimate_id:
                matches=[q for q in snapshot['estimates'] if q['kind']==expected]
                estimate_id=matches[0]['id'] if matches else None
            if not estimate_id: raise ServiceError('Сначала сохраните неизменяемую версию соответствующей сметы')
            with store.db.connection() as c: quote=get_estimate(c,estimate_id)
            if quote['repair_id']!=repair_id or quote['kind']!=expected: raise ServiceError('Выбрана другая смета или другой ремонт')
            p=quote['payload']; r=p['repair']; eq=p['equipment']; client=p['client']; organization=p['profile']; items=p['lines']; amounts=p['totals']
        values={k:str(v if v is not None else '') for k,v in r.items()}
        comparison=(quote['payload'].get('comparison') or snapshot['estimate_comparison']) if quote else snapshot['estimate_comparison']
        values.update({'repair_number':r['number'],'equipment_number':eq['number'],'equipment_model':eq['model'],'equipment_serial':eq['serial'],
                       'client_name':client['name'],'client_contact':client.get('contact',''),'client_phone':client.get('phone',''),
                       'client_email':client.get('email',''),'client_legal_name':client.get('legal_name',''),
                       'equipment_manufacturer':eq.get('manufacturer',''),'equipment_kind':eq.get('kind',''),
                       'equipment_identifiers':eq.get('identifiers',''),'status':organization.get('status_labels',{}).get(r['status'],STATUSES[r['status']][0]),
                       'organization_name':organization['organization_name'],'organization_details':organization['details'],
                       'organization_contact':organization['contact'],'warranty_terms':organization['warranty_terms'] or '____________________',
                       **{key:rub(amounts.get(key,0)) for key in ('discount','total','paid','balance')},
                       'work_total':rub(amounts['work']),'parts_total':rub(amounts['parts']),
                       'delivery_total':rub(amounts.get('delivery',0)),'other_total':rub(amounts.get('other',0)),
                       'expenses_total':rub(sum(amounts.get(k,0) for k in ('parts','delivery','other'))) if all(amounts.get(k,0) is not None for k in ('parts','delivery','other')) else 'Неизвестно',
                       'approved_preliminary_total':rub(comparison['preliminary_total']) if comparison['approved_estimate_id'] else 'не зафиксирован',
                       'estimate_difference':rub(comparison['difference']) if comparison['approved_estimate_id'] else 'не определена'})
        if quote:
            states={'prepared':'Подготовлена','sent':'Передача отмечена вручную','approved':'Согласование отмечено вручную','rejected':'Отклонение отмечено вручную','withdrawn':'Отметка отозвана'}
            values.update({'estimate_version':str(quote['version']),'estimate_date':quote['created_at'],
                           'estimate_prepared_by':quote['prepared_by'],'estimate_state':states[quote['state']],
                           'estimate_reason':quote['reason'],'estimate_comment':quote['payload']['comment']})
        for field in ('received_at','completed_at','issued_at','started_at'):
            raw_date=r.get(field,'')
            try:formatted=datetime.fromisoformat(raw_date).strftime('%d.%m.%Y') if raw_date else ''
            except ValueError:formatted=raw_date
            values[field.removesuffix('_at')+'_date']=formatted
        if quote:
            values['estimate_day']=datetime.fromisoformat(quote['created_at']).strftime('%d.%m.%Y')
        choices=image_choices if image_choices is not None else [{'id':key} for key in image_ids]
        if len({x['id'] for x in choices})!=len(choices): raise ServiceError('Одно изображение выбрано несколько раз')
        images=[]; selection=[]
        for choice in choices:
            item=store.get('attachments',choice['id'])
            if item['repair_id']!=repair_id or item['kind']!='image': raise ServiceError('Изображение должно относиться к этому ремонту')
            caption=text(choice.get('caption',item['comment'] or item['name']),'Подпись',limit=2000)
            stage=text(choice.get('stage',''),'Группа изображения',limit=60)
            images.append((attachment_path(store,item['id']),caption,stage))
            selection.append({'id':item['id'],'sha256':item['sha256'],'caption':caption,'stage':stage})
        relative_template=organization.get('template_paths',{}).get(kind)
        path=Path(template) if template else (safe_path(store.root,relative_template) if relative_template else create_templates(safe_path(store.root,'templates/default-v3'))[kind])
        logo=safe_path(store.root,organization['logo_path']) if organization.get('logo_path') else None
        for file in (path if relative_template and not template else None,logo):
            if file:
                with store.db.connection() as c: meta=c.execute('SELECT * FROM managed_files WHERE path=?',(file.relative_to(store.root).as_posix(),)).fetchone()
                if meta is None or not file.is_file() or digest(file)!=meta['sha256']: raise ServiceError('Файл профиля изменён или отсутствует')
        history=[]
        raw=Template.load(path)
        with zipfile.ZipFile(io.BytesIO(raw.raw)) as z:
            has_history=any('[[items:history]]' in ''.join(t.text or '' for t in parse(z.read(n)).iter(W+'t')) for n in z.namelist() if STORY.fullmatch(n))
        if has_history: history=_history(store,r)
        rendered=render_service(path,values,items,images,history=history,logo=logo)
        metadata={'repair_revision':r['revision'],'template_sha256':digest(path),'estimate_id':quote['id'] if quote else None,
                  'estimate_sha256':quote['payload_sha256'] if quote else None,'profile_id':organization['id'],'images':selection}
        target=None; created=False
        try:
            with store.db.connection(True) as db:
                current=db.execute('SELECT revision FROM repair_orders WHERE id=?',(repair_id,)).fetchone()
                if not quote and current[0]!=snapshot['repair']['revision']: raise ServiceError('Карточка изменилась во время генерации; повторите по свежей версии')
                version=db.execute('SELECT COALESCE(MAX(version),0)+1 FROM generated_documents WHERE repair_id=? AND kind=?',(repair_id,kind)).fetchone()[0]
                relative=f'generated/{repair_id}/{kind}-v{version:03d}-{uuid4().hex[:8]}.docx'
                target=safe_path(store.root,relative); target.parent.mkdir(parents=True,exist_ok=True)
                with target.open('xb') as f: created=True; f.write(rendered)
                db.execute('INSERT INTO generated_documents(id,repair_id,kind,version,created_at,path,sha256,template,metadata_json) VALUES(?,?,?,?,?,?,?,?,?)',
                           (uuid4().hex,repair_id,kind,version,now(),relative,digest(target),path.name,json.dumps(metadata,ensure_ascii=False)))
                store.audit(db,'document',repair_id,'generate',kind+':'+str(version))
        except Exception:
            if created: target.unlink(missing_ok=True)
            raise
        return target
