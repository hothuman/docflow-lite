"""Explicit, previewable historical import. Identity is supplied by the operator.

No fuzzy match, no inferred relationship between repair and equipment numbers.
All selected rows, source records and copied files are committed as one unit.
"""
import hashlib
import json
import shutil
import tempfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from ..data import load_table
from ..safety import check_archive, read_bounded
from .attachments import digest, inspect_file, safe_path, copy_bounded
from .estimates import canonical, normalize_lines
from .finance import calculate
from .profiles import profile
from .statuses import STATUSES
from .validation import ServiceError, day, money, now, text

HISTORY_FIELDS = {
 'repair_number':'Исторический номер ремонта *', 'repair_key':'Внешний ID ремонта (если есть)',
 'equipment_number':'Постоянный номер оборудования', 'model':'Модель', 'manufacturer':'Производитель', 'serial':'Серийный номер',
 'client_key':'Внешний ID клиента (если есть)', 'client_name':'Клиент', 'client_kind':'Тип клиента (organization / individual / entrepreneur)', 'client_contact':'Контактное лицо',
 'client_phone':'Телефон', 'client_email':'Email', 'received_at':'Дата приёмки *',
 'complaint':'Жалоба', 'defects':'Дефектовка', 'plan':'План', 'works':'Описание выполненных работ', 'parts':'Описание запчастей',
 'work_cost':'Общая стоимость работ', 'parts_cost':'Общая стоимость запчастей', 'delivery_cost':'Общая доставка',
 'other_cost':'Другие согласованные расходы', 'discount':'Скидка', 'final_total':'Окончательный итог после скидки',
 'status':'Производственный статус', 'notes':'Примечания', 'source_document':'Локальный исходный документ',
 'source_fragment':'Координаты исходного фрагмента', 'item_kind':'Тип отдельной позиции', 'item_name':'Название отдельной позиции',
 'quantity':'Количество позиции', 'unit':'Единица позиции', 'price':'Цена отдельной позиции',
}


def workbook_sheets(path):
    path = Path(path)
    if path.suffix.lower() == '.csv': return []
    if path.suffix.lower() != '.xlsx': raise ServiceError('Выберите XLSX или CSV')
    from openpyxl import load_workbook
    import io
    raw = read_bounded(path); check_archive(raw)
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try: return list(wb.sheetnames)
    finally: wb.close()


def import_day(value):
    value = text(value, 'Дата', True, 40)
    try:
        if len(value) == 10 and value[4] == '-': return day(value)
        return datetime.strptime(value, '%d.%m.%Y').date().isoformat()
    except ValueError:
        raise ServiceError('Дата: используйте ГГГГ-ММ-ДД или ДД.ММ.ГГГГ; неоднозначные даты не угадываются') from None


def import_money(value):
    return money(str(value).replace('\u00a0', '').replace(' ', ''))


def _load(path, sheet):
    path = Path(path)
    if path.is_symlink() or path.suffix.lower() not in ('.csv', '.xlsx'):
        raise ServiceError('Выберите обычный локальный XLSX/CSV')
    raw = read_bounded(path)
    # Core parses an immutable copy of the exact bytes whose hash is recorded.
    with tempfile.TemporaryDirectory(prefix='docflow-import-') as folder:
        copy = Path(folder)/('source'+path.suffix.lower()); copy.write_bytes(raw)
        table = load_table(copy, sheet=sheet)
    return table, hashlib.sha256(raw).hexdigest()


def _source_document(root, value):
    if not root: raise ServiceError('Для копирования исходных документов выберите разрешённую локальную папку')
    root = Path(root).absolute()
    p = Path(value)
    if '://' in value or str(p).startswith('\\\\'):
        raise ServiceError('Сетевые ссылки исходных документов не открываются')
    p = p if p.is_absolute() else root/p
    if not p.resolve().is_relative_to(root.resolve()): raise ServiceError('Исходный документ находится вне разрешённой папки')
    for part in (p, *p.parents):
        if part.is_symlink(): raise ServiceError('Ссылки на исходные документы запрещены')
        if part == root: break
    return p


def _groups(table, mapping, status_map, default_status):
    groups = {}; seen = set()
    for record in table.records:
        raw = {key:text(record.values[header]) for key,header in mapping.items()}
        number = text(raw.get('repair_number'), limit=100)
        key = text(raw.get('repair_key') or number, limit=300).casefold() or 'row:'+str(record.row_number)
        g = groups.setdefault(key, {'key':key, 'data':{}, 'rows':[], 'lines':[], 'errors':[], 'warnings':[]})
        g['rows'].append({'row':record.row_number, 'values':dict(record.values)})
        if not number: g['errors'].append(f'Строка {record.row_number}: нет номера ремонта')
        signature = (key, canonical(raw))
        if signature in seen: g['errors'].append(f'Строка {record.row_number}: повторяется идентичная строка; подтвердите структуру исходного журнала')
        seen.add(signature)
        for field,value in raw.items():
            if field in ('item_kind','item_name','quantity','unit','price'): continue
            if value and g['data'].get(field) and g['data'][field] != value:
                g['errors'].append(f'Строка {record.row_number}: разные значения «{HISTORY_FIELDS[field]}» у одного ремонта')
            if value: g['data'][field] = value
        if raw.get('item_name'):
            kind = raw.get('item_kind','')
            kind = {'работа':'work','запчасть':'part','доставка':'delivery','другой расход':'other'}.get(kind.casefold(),kind)
            try:
                row = normalize_lines([{'kind':kind, 'name':raw['item_name'], 'quantity':raw.get('quantity') or '1',
                                        'unit':raw.get('unit') or 'шт.', 'price':str(raw.get('price') or '0').replace(' ','').replace('\u00a0','')}])[0]
                g['lines'].append(row)
            except ServiceError as exc: g['errors'].append(f'Строка {record.row_number}: {exc}')
    for g in groups.values():
        d = g['data']
        try:
            d['received_at'] = import_day(d.get('received_at',''))
            if d.get('client_kind') and d['client_kind'] not in ('organization','individual','entrepreneur'): raise ServiceError('Неизвестный тип клиента')
            production = d.get('status','')
            if production:
                if production not in status_map: raise ServiceError('Явно сопоставьте производственный статус: '+production)
                d['status'] = status_map[production]
            else: d['status'] = default_status
            if d['status'] not in STATUSES: raise ServiceError('Выберите производственный статус для пустых исходных значений')
            for kind,description,cost in (('work','works','work_cost'),('part','parts','parts_cost'),('delivery',None,'delivery_cost'),('other',None,'other_cost')):
                if d.get(description or '') or d.get(cost):
                    if any(x['kind']==kind for x in g['lines']): raise ServiceError('Для одного типа указаны и отдельные позиции, и общий итог; выберите один формат')
                    names = {'work':'Исторические работы','part':'Исторические запчасти','delivery':'Историческая доставка','other':'Другой исторический расход'}
                    value = import_money(d.get(cost,'0'))
                    descr = d.get(description or '', '')
                    g['lines'] += normalize_lines([{'kind':kind,'name':descr.split('\n')[0][:500] if descr else names[kind],
                                                   'description':descr,'quantity':'1','unit':'компл.', 'price_cents':value}])
            discount = import_money(d.get('discount','0'))
            if d.get('final_total'):
                expected = import_money(d['final_total'])
                gross = sum(x['total_cents'] for x in g['lines'])
                if gross == 0 and expected+discount:
                    g['lines'] += normalize_lines([{'kind':'other','name':'Историческая стоимость без разбивки','price_cents':expected+discount}])
                    g['warnings'].append('Общий итог сохранён отдельно без выдуманной разбивки стоимости')
                if calculate(g['lines'],discount)['total'] != expected:
                    raise ServiceError('Сумма позиций и скидки не совпадает с указанным окончательным итогом')
            g['totals'] = calculate(g['lines'],discount); d['discount_cents'] = discount
            for position, row in enumerate(g['lines']): row['position'] = position
            if len(g['lines']) > 2000: raise ServiceError('Слишком много позиций у одного ремонта')
        except ServiceError as exc: g['errors'].append(str(exc))
        d['equipment_key'] = text(d.get('equipment_number') or 'repair:'+g['key']).casefold()
        d['client_key'] = text(d.get('client_key') or d.get('client_name') or 'repair:'+g['key']).casefold()
        if not d.get('equipment_number'): g['warnings'].append('Нет постоянного номера: аппарат нужно выбрать вручную')
        if 'client_key' not in mapping: g['warnings'].append('Нет внешнего ID клиента: одноимённых клиентов объединяет только оператор')
        semantic = {'data':d, 'lines':[{k:v for k,v in x.items() if k!='id'} for x in g['lines']]}
        g['record_sha256'] = hashlib.sha256(canonical(semantic).encode()).hexdigest()
    return list(groups.values())


def preview_history(store, path, mapping, *, sheet=None, equipment_decisions=None, client_decisions=None,
                    current_owners=None, status_map=None, default_status=None, source_label='',
                    document_root=None, copy_documents=False):
    mapping = dict(mapping); equipment_decisions = dict(equipment_decisions or {}); client_decisions = dict(client_decisions or {})
    current_owners = dict(current_owners or {}); status_map = dict(status_map or {})
    if set(mapping)-HISTORY_FIELDS.keys() or not {'repair_number','received_at'}.issubset(mapping):
        raise ServiceError('Сопоставьте номер ремонта и дату; неизвестные поля не допускаются')
    if mapping.get('equipment_number') and mapping['equipment_number'] == mapping['repair_number']:
        raise ServiceError('Номер ремонта не считается постоянным номером оборудования. Выберите отдельную колонку или свяжите аппарат вручную')
    table, file_sha = _load(path, sheet)
    if set(mapping.values())-set(table.headers): raise ServiceError('Сопоставленная колонка отсутствует в таблице')
    source_label = text(source_label, 'Идентификатор источника', limit=300)
    source_id = hashlib.sha256(canonical([source_label or file_sha,sheet or '']).encode()).hexdigest()
    groups = _groups(table,mapping,status_map,default_status)
    with store.db.connection() as c:
        equipment = {r['id']:dict(r) for r in c.execute('SELECT * FROM equipment')}
        clients = {r['id']:dict(r) for r in c.execute('SELECT * FROM clients')}
        existing = {r['source_key']:dict(r) for r in c.execute('SELECT * FROM import_records WHERE source_id=?',(source_id,))}
        repairs = {r['number'].casefold():dict(r) for r in c.execute('SELECT id,number,equipment_id,client_id FROM repair_orders')}
    owners = defaultdict(set)
    for g in groups: owners[g['data']['equipment_key']].add(g['data']['client_key'])
    for g in groups:
        d = g['data']; eqkey = d['equipment_key']; ckey = d['client_key']
        g['equipment_candidates'] = [x['id'] for x in equipment.values() if x['number'].casefold() == d.get('equipment_number','').casefold()]
        g['client_candidates'] = [x['id'] for x in clients.values() if x['name'].casefold() == d.get('client_name','').casefold()]
        g['equipment_decision'] = equipment_decisions.get(eqkey)
        g['client_decision'] = client_decisions.get(ckey)
        if g['key'] in existing:
            g['status'] = 'already_imported' if existing[g['key']]['record_sha256']==g['record_sha256'] else 'conflict'
            if g['status']=='conflict': g['errors'].append('Такая запись источника уже импортирована с другим содержимым; автоматического обновления нет')
            continue
        if g['equipment_decision']=='defer' or g['client_decision']=='defer':
            g['status']='deferred'; continue
        if d.get('repair_number','').casefold() in repairs:
            g['errors'].append('Номер ремонта уже существует в базе из другого источника; нужен ручной разбор')
        if g['equipment_decision']=='new':
            if not d.get('equipment_number') or not d.get('model'): g['errors'].append('Для нового аппарата нужны постоянный номер и модель')
            if g['equipment_candidates']: g['errors'].append('Такой постоянный номер уже занят; выберите существующий аппарат')
            if len(owners[eqkey])>1 and not current_owners.get(eqkey): g['errors'].append('Несколько исторических заказчиков: явно выберите текущего владельца нового аппарата')
        elif g['equipment_decision']:
            eq = equipment.get(g['equipment_decision'])
            if eq is None or eq['archived']: g['errors'].append('Выбранный аппарат отсутствует или архивирован')
        if g['client_decision']=='new':
            if not d.get('client_name'): g['errors'].append('Для нового клиента нужно имя/наименование')
            email = d.get('client_email','')
            if email and ('@' not in email or any(x.isspace() for x in email)): g['errors'].append('Неверный email клиента')
        elif g['client_decision']:
            cl = clients.get(g['client_decision'])
            if cl is None or cl['archived']: g['errors'].append('Выбранный клиент отсутствует или архивирован')
        if copy_documents and d.get('source_document'):
            try:
                p = _source_document(document_root,d['source_document']); inspect_file(p)
            except (ServiceError,OSError) as exc: g['errors'].append('Исходный документ: '+str(exc))
        elif d.get('source_document'):
            g['warnings'].append('Сохранена ссылка-источник; файл не копируется и автоматически не открывается')
        g['status'] = 'error' if g['errors'] else ('ready' if g['equipment_decision'] and g['client_decision'] else 'needs_confirmation')
    options = dict(sheet=sheet,equipment_decisions=equipment_decisions,client_decisions=client_decisions,current_owners=current_owners,
                   status_map=status_map,default_status=default_status,source_label=source_label,
                   document_root=str(document_root) if document_root else None,copy_documents=bool(copy_documents))
    return {'path':str(Path(path).absolute()),'file_sha256':file_sha,'source_id':source_id,'mapping':mapping,'options':options,
            'headers':list(table.headers),'groups':groups,'records_previewed':len(table.records), 'writes_performed':False}


def import_history(store, plan, *, confirmed=False, selected_keys=None):
    if not confirmed: raise ServiceError('Импорт требует подтверждения сопоставления и выбранных строк')
    if digest(Path(plan['path'])) != plan['file_sha256']:
        raise ServiceError('Исходный файл изменился после предпросмотра; проверьте его заново')
    created = []
    try:
        with store.lock, store.db.connection(True) as c:
            fresh = preview_history(store,plan['path'],plan['mapping'],**plan['options'])
            selected = set(selected_keys) if selected_keys is not None else {g['key'] for g in fresh['groups'] if g['status']!='deferred'}
            known = {g['key'] for g in fresh['groups']}
            if selected-known: raise ServiceError('В подтверждении есть неизвестные строки')
            rows = [g for g in fresh['groups'] if g['key'] in selected]
            if any(g['status'] not in ('ready','already_imported') for g in rows):
                raise ServiceError('Есть ошибки/неподтверждённые строки; транзакция отменена. Отложите их явно или исправьте источник')
            old = {g['key']:g['record_sha256'] for g in plan['groups']}
            if any(old.get(g['key'])!=g['record_sha256'] for g in rows): raise ServiceError('План импорта изменился; нужен новый просмотр')
            if not rows: raise ServiceError('Не выбраны подтверждённые строки')
            stamp = now(); source = fresh['source_id']
            c.execute('INSERT OR IGNORE INTO import_sources(id,name,file_sha256,sheet,mapping_json,created_at) VALUES(?,?,?,?,?,?)',
                      (source,Path(fresh['path']).name,fresh['file_sha256'],fresh['options']['sheet'] or '',canonical(fresh['mapping']),stamp))
            client_cache = {}; equipment_cache = {}
            # An explicitly chosen current owner may belong to a deferred old
            # repair. Create only that confirmed client, not the deferred repair.
            by_client = {g['data']['client_key']:g for g in fresh['groups']}
            def client(key):
                if key in client_cache: return client_cache[key]
                if key not in by_client:
                    record = store.require(c,'clients',key)
                    if record['archived']: raise ServiceError('Владелец архивирован')
                    return key
                g = by_client[key]; decision = g['client_decision']; d = g['data']
                if not decision or decision=='defer':raise ServiceError('Текущий владелец не подтверждён в сопоставлении клиентов')
                if decision!='new':
                    record=store.require(c,'clients',decision)
                    if record['archived']: raise ServiceError('Клиент архивирован')
                    client_cache[key]=decision; return decision
                ident=uuid4().hex
                email=text(d.get('client_email'),limit=300)
                if email and ('@' not in email or any(x.isspace() for x in email)):raise ServiceError('Неверный email выбранного владельца')
                c.execute('INSERT INTO clients(id,kind,name,contact,phone,email,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)',
                          (ident,d.get('client_kind') or 'organization',text(d.get('client_name'),'Клиент',True,300),text(d.get('client_contact'),limit=300),
                           text(d.get('client_phone'),limit=300),email,stamp,stamp))
                client_cache[key]=ident; return ident
            report=[]
            for g in rows:
                if g['status']=='already_imported': report.append({'key':g['key'],'rows':g['rows'],'result':'already_imported'}); continue
                d=g['data']; customer=client(d['client_key']); eqkey=d['equipment_key']; decision=g['equipment_decision']
                if eqkey not in equipment_cache:
                    if decision=='new':
                        current_owner=fresh['options']['current_owners'].get(eqkey) or d['client_key']
                        owner=client(current_owner); eq=uuid4().hex
                        if c.execute('SELECT 1 FROM equipment WHERE lower(number)=lower(?)',(d['equipment_number'],)).fetchone():
                            raise ServiceError('Постоянный номер занят после предпросмотра')
                        c.execute('INSERT INTO equipment(id,number,model,manufacturer,serial,owner_id,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)',
                                  (eq,text(d['equipment_number'],'Номер',True,300),text(d['model'],'Модель',True,300),
                                   text(d.get('manufacturer'),limit=300),text(d.get('serial'),limit=300),owner,stamp,stamp))
                        equipment_cache[eqkey]=eq
                    else:
                        record=store.require(c,'equipment',decision)
                        if record['archived']: raise ServiceError('Аппарат архивирован')
                        equipment_cache[eqkey]=decision
                repair=uuid4().hex; active=profile(c,active=True)
                c.execute('INSERT INTO repair_orders(id,number,equipment_id,client_id,received_at,complaint,defects,plan,notes,status,discount_cents,created_at,updated_at,profile_id) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                          (repair,text(d['repair_number'],'Номер ремонта',True,100),equipment_cache[eqkey],customer,d['received_at'],
                           text(d.get('complaint')),text(d.get('defects')),text(d.get('plan')),text(d.get('notes')),d['status'],d['discount_cents'],stamp,stamp,active['id']))
                for item in g['lines']:
                    c.execute('INSERT INTO repair_items(id,repair_id,kind,name,description,quantity,unit,price_cents,total_cents,position,notes) VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                              (uuid4().hex,repair,item['kind'],item['name'],item['description'],item['quantity'],item['unit'],item['price_cents'],item['total_cents'],item['position'],item['notes']))
                c.execute('INSERT INTO repair_status_history(repair_id,old_status,new_status,changed_at,note,source) VALUES(?,?,?,?,?,?)',
                          (repair,'',d['status'],stamp,'Историческая запись; дата приёмки сохранена отдельно','import'))
                if fresh['options']['copy_documents'] and d.get('source_document'):
                    src=_source_document(fresh['options']['document_root'],d['source_document']); kind,mime=inspect_file(src)
                    ident=uuid4().hex; relative=f'attachments/{equipment_cache[eqkey]}/{repair}/{ident}{src.suffix.lower()}'
                    target=safe_path(store.root,relative); target.parent.mkdir(parents=True,exist_ok=True)
                    copy_bounded(src,target); created.append(target)
                    inspect_file(target)
                    c.execute('INSERT INTO attachments(id,repair_id,kind,name,path,size,mime,sha256,comment,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)',
                              (ident,repair,kind,text(src.name,limit=255),relative,target.stat().st_size,mime,digest(target),'Исходный документ импорта',stamp))
                c.execute('INSERT INTO import_records(id,source_id,source_key,record_sha256,repair_id,rows_json,imported_at) VALUES(?,?,?,?,?,?,?)',
                          (uuid4().hex,source,g['key'],g['record_sha256'],repair,canonical({'file_sha256':fresh['file_sha256'],'rows':g['rows'],'source_document':d.get('source_document',''),'fragment':d.get('source_fragment','')}),stamp))
                store.audit(c,'repair',repair,'import',source+':'+g['key'])
                report.append({'key':g['key'],'rows':g['rows'],'repair_id':repair,'result':'imported','warnings':g['warnings']})
            for g in fresh['groups']:
                if g['key'] not in selected: report.append({'key':g['key'],'rows':g['rows'],'result':'deferred','issues':g['errors']})
            result={'status':'completed' if any(r['result']=='imported' for r in report) else 'no_changes',
                    'source_id':source,'file_sha256':fresh['file_sha256'],'created_at':stamp,'rows':report,
                    'imported':sum(r['result']=='imported' for r in report),'already_imported':sum(r['result']=='already_imported' for r in report),
                    'deferred':sum(r['result']=='deferred' for r in report),'transaction':'all selected changes committed together'}
            run=uuid4().hex
            c.execute('INSERT INTO import_runs(id,source_id,report_json,created_at) VALUES(?,?,?,?)',(run,source,canonical(result),stamp))
            result['run_id']=run
        return result
    except Exception:
        for path in created: path.unlink(missing_ok=True)
        raise
