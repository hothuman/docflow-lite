"""Review irregular workbooks without guessing entities or importing records.

Specs explicitly provide each sheet's start row and semantic column letters.
Reports retain source values and must remain in the private customer folder.
"""
from datetime import date, datetime
import csv
import hashlib
import io
import re
from pathlib import Path

from ..safety import check_archive, read_bounded
from .validation import ServiceError, money

MONEY = re.compile(r'\+?(?:\d{1,3}(?:[ \u00a0]\d{3})+|\d+)(?:[.,]\d{1,2}|-\d{2})?\s*(?:руб\.?|₽)?', re.I)
DATE = re.compile(r'(?<!\d)(\d{1,2})[./](\d{1,2})[./](\d{4})(?!\d)')


def legacy_money(value):
    """Absent is None, never zero. Arithmetic and alternative totals are rejected."""
    if value is None or str(value).strip() == '':
        return None
    if isinstance(value, bool):
        raise ServiceError('Логическое значение не является суммой')
    raw = str(value).strip()
    if not MONEY.fullmatch(raw):
        raise ServiceError('Сумма неоднозначна; требуется проверенное значение')
    raw = re.sub(r'\s*(?:руб\.?|₽)$', '', raw, flags=re.I)
    return money(raw.replace(' ', '').replace('\u00a0', '').replace('-', '.'))


def legacy_date_number(value):
    result = {'received_at': '', 'repair_number': '', 'issues': []}
    if isinstance(value, (date, datetime)):
        result['received_at'] = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    else:
        raw = str(value or '').strip()
        matches = list(DATE.finditer(raw))
        if len(matches) == 1:
            m = matches[0]
            try:
                result['received_at'] = date(int(m[3]), int(m[2]), int(m[1])).isoformat()
            except ValueError:
                result['issues'].append('invalid_date')
            remainder = (raw[:m.start()] + raw[m.end():]).strip(' \r\n№#;,')
            if re.fullmatch(r'\d+(?:[-/]\d+)*', remainder):
                result['repair_number'] = remainder
            elif remainder:
                result['issues'].append('ambiguous_repair_number')
        elif re.fullmatch(r'\d{4}-\d{2}-\d{2}', raw):
            try:
                result['received_at'] = date.fromisoformat(raw).isoformat()
            except ValueError:
                result['issues'].append('invalid_date')
        else:
            result['issues'].append('ambiguous_date' if raw else 'missing_date')
    if not result['repair_number']:
        result['issues'].append('missing_repair_number')
    return result


def review_workbook(path, sheet_specs):
    """Read all physical rows. Explicitly exclude any irrelevant sheet with a reason."""
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string
    path = Path(path)
    if path.suffix.lower() != '.xlsx' or path.is_symlink():
        raise ServiceError('Выберите обычный локальный XLSX')
    raw = read_bounded(path); check_archive(raw)
    sha = hashlib.sha256(raw).hexdigest()
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        if set(wb.sheetnames) != set(sheet_specs):
            raise ServiceError('Явно задайте схему каждого листа')
        sheets = []
        for ws in wb:
            spec = sheet_specs[ws.title]
            if spec.get('exclude'):
                if not spec.get('reason'):
                    raise ServiceError('Укажите причину исключения листа')
                sheets.append({'sheet': ws.title, 'excluded': True, 'reason': spec['reason'], 'rows': []})
                continue
            start = spec.get('start_row')
            if type(start) is not int or start < 1:
                raise ServiceError('Задайте первую строку данных')
            if ws.max_row > 100000 or ws.max_column > 1000:
                raise ServiceError('Слишком большой журнал для построчной проверки')
            mapping = {k: column_index_from_string(v) for k, v in spec['columns'].items()}
            rows = []
            for row, cells in enumerate(ws.iter_rows(min_row=start), start):
                original = {c.coordinate: {'value': c.value.isoformat() if isinstance(c.value, (date, datetime)) else c.value,
                                          'type': c.data_type, 'number_format': c.number_format}
                            for c in cells if c.value is not None}
                if not original:
                    continue
                values = {k: cells[n-1].value if n <= len(cells) else None for k, n in mapping.items()}
                parsed = legacy_date_number(values.get('date_number'))
                issues = parsed.pop('issues')
                costs = {}
                for field in ('work_cost', 'parts_cost', 'final_total'):
                    try:
                        costs[field] = legacy_money(values.get(field))
                    except ServiceError:
                        costs[field] = None
                        issues.append('ambiguous_' + field)
                if costs['final_total'] is None:
                    issues.append('unknown_total')
                if all(costs[k] is not None for k in ('work_cost', 'parts_cost', 'final_total')):
                    if costs['work_cost'] + costs['parts_cost'] != costs['final_total']:
                        issues.append('cost_mismatch')
                if not values.get('date_number'):
                    issues.append('possible_continuation')
                if spec.get('combined_identity'):
                    issues.append('combined_identity')
                if any(c['type'] == 'f' for c in original.values()):
                    issues.append('formula_requires_review')
                issues += ['equipment_identity_requires_confirmation', 'client_identity_requires_confirmation',
                           'production_status_requires_confirmation', 'payment_unknown']
                rows.append({'source_key': hashlib.sha256(f'{sha}\0{ws.title}\0{row}'.encode()).hexdigest(),
                             'row': row, 'cells': original, 'mapped_raw': values,
                             'proposed': {**parsed, 'costs_cents': costs}, 'issues': sorted(set(issues)),
                             'approved': False})
            sheets.append({'sheet': ws.title, 'dimension': ws.calculate_dimension(), 'start_row': start, 'rows': rows})
        return {'file_name': path.name, 'file_sha256': sha, 'sheets': sheets,
                'physical_rows': sum(len(s['rows']) for s in sheets), 'writes_performed': False,
                'scope': 'Physical rows, not confirmed repairs. No entity merging, payment inference or import.'}
    finally:
        wb.close()


def export_review_csv(report, target):
    def safe(value):
        value = str(value if value is not None else '')
        return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value
    with Path(target).open('x', encoding='utf-8-sig', newline='') as f:
        out = csv.writer(f, delimiter=';')
        out.writerow(['source_key', 'Лист', 'Строка', 'Предложенная дата', 'Предложенный номер обращения',
                      'Итог из источника, коп.', 'Проверить', 'Решение', 'Постоянный номер аппарата',
                      'ID клиента', 'Проверенная дата', 'Проверенный номер', 'Примечание проверяющего'])
        for sheet in report['sheets']:
            for row in sheet['rows']:
                p = row['proposed']
                out.writerow([safe(x) for x in [row['source_key'], sheet['sheet'], row['row'], p['received_at'],
                              p['repair_number'], p['costs_cents']['final_total'], ', '.join(row['issues']),
                              '', '', '', '', '', '']])

def prepare_history_rows(report, decisions):
    """Convert only explicitly approved groups to the existing importer's fields.

    Decisions are separate from machine proposals. Unmentioned rows stay open;
    no source row can participate in two repairs. This function never writes DB.
    """
    import json
    from .history_import import HISTORY_FIELDS
    source={r['source_key']:(sheet['sheet'],r) for sheet in report['sheets'] for r in sheet['rows']}
    used=set();result=[];deferred=[]
    for decision in decisions:
        keys=decision.get('source_keys')
        if not isinstance(keys,list) or not keys or len(keys)!=len(set(keys)):
            raise ServiceError('Решение должно перечислять неповторяющиеся исходные строки')
        if set(keys)-source.keys() or set(keys)&used:
            raise ServiceError('Неизвестная или повторно использованная исходная строка')
        action=decision.get('action')
        if action not in ('approve','defer'):
            raise ServiceError('Явно выберите approve или defer')
        used.update(keys)
        if action=='defer':
            deferred.extend(keys);continue
        actor=str(decision.get('reviewed_by') or '').strip()
        note=str(decision.get('reason') or '').strip()
        if not actor or not note:
            raise ServiceError('Укажите проверяющего и основание решения')
        data=dict(decision.get('data') or {})
        if set(data)-HISTORY_FIELDS.keys():
            raise ServiceError('Неизвестные поля нормализованной истории')
        required={'repair_number','received_at','equipment_number','model','client_key','client_name'}
        if any(not str(data.get(k) or '').strip() for k in required):
            raise ServiceError('Подтвердите номер/дату обращения, аппарат и отдельный ID клиента')
        ordered=sorted(keys)
        data['repair_key']=hashlib.sha256(json.dumps(ordered).encode()).hexdigest()
        data['source_fragment']=json.dumps({'file_sha256':report['file_sha256'],
            'rows':[{'sheet':source[k][0],'row':source[k][1]['row'],'source_key':k} for k in ordered],
            'reviewed_by':actor,'reason':note},ensure_ascii=False)
        result.append(data)
    return {'rows':result,'approved_source_rows':sum(len(d['source_keys']) for d in decisions if d.get('action')=='approve'),
            'deferred_source_rows':len(deferred),'unreviewed_source_rows':len(source)-len(used),
            'scope':'Approved normalized rows for history_import; no database writes.'}
