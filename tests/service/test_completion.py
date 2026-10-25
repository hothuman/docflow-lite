"""Completion regressions on real SQLite, files and DOCX; no mock repository."""
import json
import sqlite3
from contextlib import closing

import pytest
from docx import Document

from docflow.service.database import Database, MIGRATIONS, SCHEMA_VERSION
from docflow.service.repository import ServiceStore
from docflow.service.demo import seed_demo
from docflow.service.documents import generate_document
from docflow.service.backup import create_backup, restore_backup
from docflow.service.attachments import digest
from docflow.service.validation import ServiceError
from docflow.service.finance import totals


def legacy_database(root):
    root.mkdir()
    path = root / 'service.db'
    doc = root / 'generated/legacy.docx'
    doc.parent.mkdir()
    d = Document(); d.add_paragraph('Исторический синтетический документ'); d.save(doc)
    with closing(sqlite3.connect(path)) as c:
        for version in (1, 2):
            for sql in MIGRATIONS[version]: c.execute(sql)
        c.execute("INSERT INTO clients(id,kind,name,created_at,updated_at) VALUES('c','organization','Учебный клиент','2023','2023')")
        c.execute("INSERT INTO equipment(id,number,model,owner_id,created_at,updated_at) VALUES('e','900','Учебный аппарат','c','2023','2023')")
        c.execute("INSERT INTO repair_orders(id,number,equipment_id,client_id,received_at,created_at,updated_at) VALUES('r','101','e','c','2023-02-01','2023','2023')")
        c.execute("INSERT INTO repair_items(id,repair_id,kind,name,quantity,unit,price_cents,total_cents) VALUES('i','r','work','Работа','1.5','ч',100,150)")
        c.execute("INSERT INTO payments(id,repair_id,amount_cents,occurred_at,method,kind,created_at) VALUES('p','r',100,'2023-02-01','Учебная запись','payment','2023')")
        c.execute("INSERT INTO repair_status_history(repair_id,old_status,new_status,changed_at,source) VALUES('r','','accepted','2023','create')")
        c.execute("INSERT INTO generated_documents(id,repair_id,kind,version,created_at,path,sha256,template) VALUES('d','r','intake',1,'2023','generated/legacy.docx',?,'legacy.docx')", (digest(doc),))
        c.execute('PRAGMA user_version=2'); c.commit()
        before = {t: c.execute('SELECT * FROM '+t).fetchall() for t in ('clients','equipment','repair_orders','repair_items','payments','repair_status_history','generated_documents')}
    return path, before, doc.read_bytes()


def test_alpha_v2_migration_preserves_history_finance_files(tmp_path):
    path, before, doc = legacy_database(tmp_path/'alpha')
    db = Database(path)
    with db.connection() as c:
        assert c.execute('PRAGMA user_version').fetchone()[0] == SCHEMA_VERSION
        assert c.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
        assert not c.execute('PRAGMA foreign_key_check').fetchall()
        for table, expected in before.items():
            rows = [tuple(row)[:len(expected[0])] for row in c.execute('SELECT * FROM '+table)]
            assert rows == expected, table
        assert totals(c, 'r')['balance'] == 50
    assert (path.parent/'generated/legacy.docx').read_bytes() == doc
    copies = list(path.parent.glob('service.db.before-v3-*.bak')); assert len(copies) == 1
    with closing(sqlite3.connect(copies[0])) as c:
        assert c.execute('PRAGMA user_version').fetchone()[0] == 2
        assert c.execute('SELECT number FROM repair_orders').fetchone()[0] == '101'


def test_failed_v3_migration_atomic_and_unique_backups(tmp_path, monkeypatch):
    path, before, _ = legacy_database(tmp_path/'alpha')
    monkeypatch.setitem(MIGRATIONS, 3, ['CREATE TABLE transient(x)', 'BROKEN SQL'])
    for _ in range(2):
        with pytest.raises(ServiceError): Database(path)
    with closing(sqlite3.connect(path)) as c:
        assert c.execute('PRAGMA user_version').fetchone()[0] == 2
        assert c.execute('SELECT * FROM repair_items').fetchall() == before['repair_items']
        assert not c.execute("SELECT 1 FROM sqlite_master WHERE name='transient'").fetchone()
    assert len(list(path.parent.glob('service.db.before-v3-*.bak'))) == 2


def test_docx_commit_failure_removes_own_output(tmp_path, monkeypatch):
    s = ServiceStore(tmp_path/'project'); ids = seed_demo(s.root)
    original = sqlite3.connect
    class Full(sqlite3.Connection):
        def commit(self): raise sqlite3.OperationalError('database or disk is full')
    monkeypatch.setattr(sqlite3, 'connect', lambda *a, **kw: original(*a, **kw, factory=Full))
    with pytest.raises(ServiceError): generate_document(s, ids['repair'], 'completion')
    assert not s.snapshot(ids['repair'])['generated_documents']
    assert not list((s.root/'generated').rglob('*.docx'))


def test_custom_template_survives_backup_restore(tmp_path):
    s = ServiceStore(tmp_path/'project'); seed_demo(s.root)
    template = s.root/'templates/custom/own.docx'; template.parent.mkdir(parents=True)
    d = Document(); d.add_paragraph('Своя форма {{repair_number}}'); d.save(template)
    backup = create_backup(s, tmp_path/'copy.zip')
    restored = restore_backup(backup, tmp_path/'restored')
    assert (restored/'templates/custom/own.docx').read_bytes() == template.read_bytes()


def test_new_expense_categories_are_distinct(tmp_path):
    s = ServiceStore(tmp_path/'project'); ids = seed_demo(s.root)
    with s.db.connection(True) as c:
        for kind in ('delivery', 'other'):
            c.execute('INSERT INTO repair_items(id,repair_id,kind,name,quantity,unit,price_cents,total_cents) VALUES(?,?,?,?,?,?,?,?)',
                      (kind, ids['repair'], kind, kind, '1', 'шт', 125, 125))
    t = s.snapshot(ids['repair'])['totals']
    assert t['delivery'] == t['other'] == 125
    assert t['total'] == 270300 and t['balance'] == 170300
