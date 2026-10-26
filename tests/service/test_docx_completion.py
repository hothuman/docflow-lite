import io,json,zipfile
from pathlib import Path
import pytest
from docx import Document
from docx.enum.section import WD_SECTION,WD_ORIENT
from docx.shared import Inches
from docflow.service.demo import seed_demo
from docflow.service.repository import ServiceStore
from docflow.service.documents import generate_document,create_templates,render_service,KINDS
from docflow.service.estimates import save_draft,freeze_estimate,record_event
from docflow.service.profiles import save_profile
from docflow.service.attachments import add_attachment
from docflow.service.validation import ServiceError

def contents(p):
 d=Document(p);return '\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])

def test_templates_all_load_and_no_unsafe_variable_breaks(tmp_path):
 from docflow.templates import Template
 for path in create_templates(tmp_path).values():assert Template.load(path).fields
 assert len(KINDS)==8

def test_frozen_quote_docx_never_replaces_old_prices_comparison_or_identity(tmp_path):
 s=ServiceStore(tmp_path/'store');ids=seed_demo(s.root);r=s.create_repair(ids['equipment'],'FROZEN-001')
 save_draft(s,r,[{'kind':'work','name':'План А','price':'100'}]);q=freeze_estimate(s,r,'preliminary','Учебный оператор')
 record_event(s,q,'approved','Оператор','Тестовая отметка')
 s.save_item(r,{'kind':'other','name':'Факт А','price':'80','change_reason':'Тестовое уточнение'})
 final=freeze_estimate(s,r,'final','Оператор','Тестовое уточнение')
 before=contents(generate_document(s,r,'estimate_final',estimate_id=final))
 s.save_item(r,{'kind':'delivery','name':'Новая доставка','price':'777','change_reason':'Другой состав'})
 save_draft(s,r,[{'kind':'work','name':'Новый план','price':'999'}]);p2=freeze_estimate(s,r,'preliminary','Оператор','Новая версия')
 record_event(s,p2,'approved','Оператор','Новая тестовая отметка')
 after=contents(generate_document(s,r,'estimate_final',estimate_id=final))
 assert before==after and 'Новая доставка' not in after and '80,00 ₽' in after and '−20,00 ₽' in after
 assert 'План А' in contents(generate_document(s,r,'estimate_preliminary',estimate_id=q))
 with pytest.raises(ServiceError):generate_document(s,r,'estimate_final',estimate_id=q)

def test_complex_template_stories_history_sections_logo_and_selected_images(tmp_path):
 from PySide6.QtGui import QImage,QColor
 s=ServiceStore(tmp_path/'store');ids=seed_demo(s.root);r=ids['repair'];images=[]
 for i,(w,h) in enumerate([(40,800),(900,100),(100,100)]):
  p=tmp_path/f'image-{i}.png';im=QImage(w,h,QImage.Format.Format_RGB32);im.fill(QColor('red'));assert im.save(str(p));images.append(add_attachment(s,r,p,f'Схема {i}'))
 d=Document();d.add_paragraph('[[logo]]');d.add_paragraph('{{repair_number}}');d.sections[0].header.paragraphs[0].text='Заголовок {{client_name}}';d.sections[0].footer.paragraphs[0].text='Низ {{equipment_number}}'
 outer=d.add_table(rows=1,cols=1);t=outer.cell(0,0).add_table(rows=2,cols=3);t.cell(0,0).merge(t.cell(0,2)).text='История'
 t.cell(1,0).text='[[items:history]]{{history_number}}';t.cell(1,1).text='{{history_client}}';t.cell(1,2).text='{{history_total}}'
 section=d.add_section(WD_SECTION.NEW_PAGE);section.orientation=WD_ORIENT.LANDSCAPE;section.page_width=Inches(11.7);section.page_height=Inches(8.3)
 d.add_paragraph('Результат {{defects}}');p=tmp_path/'complex.docx';d.save(p)
 choices=[{'id':images[1],'caption':'Первая выбранная схема','stage':'После ремонта'},{'id':images[0],'caption':'Вторая выбранная схема','stage':'До ремонта'}]
 out=generate_document(s,r,'combined',p,image_choices=choices);result=Document(out)
 assert len(result.inline_shapes)==2 and len(result.sections)==2
 assert 'DEMO-R-001' in result.tables[0].cell(0,0).tables[0].cell(1,0).text
 assert '{{' not in result.sections[0].header.paragraphs[0].text
 captions=[p.text for p in result.paragraphs if 'выбранная схема' in p.text];assert captions==['Первая выбранная схема','Вторая выбранная схема']
 for index,paragraph in enumerate(result.paragraphs):
  if 'выбранная схема' not in paragraph.text:continue
  assert paragraph.paragraph_format.keep_with_next is True
  picture=result.paragraphs[index+1]
  assert picture._p.xpath('.//w:drawing')
  assert picture.paragraph_format.keep_together is True
  assert picture.paragraph_format.line_spacing==1
 for shape in result.inline_shapes:assert shape.width<=result.sections[-1].page_width-result.sections[-1].left_margin-result.sections[-1].right_margin
 meta=json.loads(s.snapshot(r)['generated_documents'][0]['metadata_json']);assert [x['id'] for x in meta['images']]==[images[1],images[0]]
 with pytest.raises(ServiceError):generate_document(s,r,'combined',image_choices=[{'id':images[0]},{'id':images[0]}])


def test_neutral_warranty_no_invented_duration_or_signatures_and_no_video(tmp_path):
 s=ServiceStore(tmp_path/'p');ids=seed_demo(s.root);p=generate_document(s,ids['repair'],'warranty');t=contents(p)
 assert 'Прототип для утверждения' in t and 'Срок и начало действия: ____________________' in t
 assert '12 месяцев' not in t and not Document(p).inline_shapes
 v=add_attachment(s,ids['repair'],Path(__file__).parents[2]/'site/dist/assets/demo.mp4')
 with pytest.raises(ServiceError):generate_document(s,ids['repair'],'warranty',image_ids=[v])


def test_expense_tables_present_only_for_real_categories(tmp_path):
 s=ServiceStore(tmp_path/'p');ids=seed_demo(s.root);r=s.create_repair(ids['equipment'])
 assert len(Document(generate_document(s,r,'completion')).tables)==2
 s.save_item(r,{'kind':'delivery','name':'Доставка','price':'20'})
 s.save_item(r,{'kind':'other','name':'Согласованный расход','price':'5'})
 d=Document(generate_document(s,r,'completion'));assert len(d.tables)==4
 assert d.tables[2].cell(1,1).text.startswith('Доставка') and '25,00 ₽' in contents(generate_document(s,r,'completion'))

def test_images_work_without_english_heading_styles_in_customer_template(tmp_path):
 from PySide6.QtGui import QImage,QColor
 from docx.enum.style import WD_STYLE_TYPE
 p=tmp_path/'template.docx';d=Document();d.add_paragraph('{{repair_number}}')
 for name in ('Heading 1','Heading 2'):
  element=d.styles[name]._element;element.getparent().remove(element)
 d.styles.add_style('Исходный стиль',WD_STYLE_TYPE.PARAGRAPH)
 d.add_paragraph('Неизменённый текст','Исходный стиль');d.save(p)
 image=tmp_path/'photo.png';im=QImage(200,100,QImage.Format.Format_RGB32);im.fill(QColor('green'));assert im.save(str(image))
 raw=render_service(p,{'repair_number':'SYNTHETIC-1'},[],[(image,'Учебная подпись','До ремонта')])
 result=Document(io.BytesIO(raw))
 assert len(result.inline_shapes)==1
 assert 'Исходный стиль' in [style.name for style in result.styles]
 assert 'Учебная подпись' in '\\n'.join(p.text for p in result.paragraphs)
 assert result.paragraphs[1].text=='Неизменённый текст'

def test_optional_global_item_index_keeps_per_group_numbering_compatible(tmp_path):
 d=Document();t=d.add_table(rows=1,cols=2);t.cell(0,0).text='[[items:work]]{{item_no}} / {{item_index}}';t.cell(0,1).text='{{item_name}}'
 t=d.add_table(rows=1,cols=2);t.cell(0,0).text='[[items:part]]{{item_no}} / {{item_index}}';t.cell(0,1).text='{{item_name}}'
 p=tmp_path/'template.docx';d.save(p)
 from docflow.service.estimates import normalize_lines
 lines=normalize_lines([{'kind':'work','name':'Work','price':'1'},{'kind':'part','name':'Part','price':'2'}])
 result=Document(io.BytesIO(render_service(p,{},lines,[])))
 assert result.tables[0].cell(0,0).text=='1 / 1'
 assert result.tables[1].cell(0,0).text=='1 / 2'
