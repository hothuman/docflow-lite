import csv
from pathlib import Path

import pytest

from docflow.service.repository import ServiceStore
from docflow.service.history_import import preview_history,import_history
from docflow.service.backup import create_backup,restore_backup
from docflow.service.estimates import freeze_estimate
from docflow.service.validation import ServiceError,rub


def imported(tmp_path,extra=None):
    extra=extra or {}
    fields={'repair_number':'R-1','equipment_number':'E-1','model':'Synthetic pump','client_key':'C-1',
            'client_name':'Synthetic customer','received_at':'2020-01-01','works':'Historical work',**extra}
    tmp_path.mkdir(parents=True,exist_ok=True)
    p=tmp_path/'source.csv'
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        out=csv.writer(f,delimiter=';');out.writerow(fields);out.writerow(fields.values())
    s=ServiceStore(tmp_path/'project')
    plan=preview_history(s,p,{k:k for k in fields},equipment_decisions={'e-1':'new'},
                         client_decisions={'c-1':'new'},default_status='closed')
    report=import_history(s,plan,confirmed=True)
    return s,report['rows'][0]['repair_id'],plan


def test_missing_payment_is_not_debt_and_replay_restore_preserve_unknown(tmp_path):
    s,r,plan=imported(tmp_path,{'final_total':'1250-50'})
    t=s.snapshot(r)['totals']
    assert t['total']==125050 and t['paid'] is None and t['balance'] is None
    assert t['recorded_paid']==0 and 'неизвест' in t['financial_status'].lower()
    assert s.dashboard()['debt']==0
    assert s.repairs()[0]['balance'] is None
    assert import_history(s,plan,confirmed=True)['already_imported']==1
    create_backup(s,tmp_path/'copy.zip');restore_backup(tmp_path/'copy.zip',tmp_path/'restored')
    assert ServiceStore(tmp_path/'restored').snapshot(r)['totals']==t
    assert rub(None)=='Неизвестно'


def test_missing_cost_is_not_a_free_repair_but_explicit_zero_is_zero(tmp_path):
    s,r,_=imported(tmp_path)
    assert s.snapshot(r)['totals']['total'] is None
    with pytest.raises(ServiceError):freeze_estimate(s,r,'final','Synthetic operator')
    s2,r2,_=imported(tmp_path/'zero',{'final_total':'0'})
    assert s2.snapshot(r2)['totals']['total']==0


def test_explicit_payment_import_and_controlled_confirmation(tmp_path):
    s,r,_=imported(tmp_path,{'final_total':'1500','paid_amount':'600','payment_state':'known',
                            'payment_date':'2020-01-02','payment_method':'Historical receipt'})
    t=s.snapshot(r)['totals']
    assert (t['total'],t['paid'],t['balance'])==(150000,60000,90000)
    assert s.snapshot(r)['payments'][0]['occurred_at']=='2020-01-02'
    rev=s.get('repair_orders',r)['revision']
    with pytest.raises(ServiceError):s.set_financial_evidence(r,True,False,'',expected_revision=rev)
    s.set_financial_evidence(r,True,False,'Payment history incomplete',expected_revision=rev)
    assert s.snapshot(r)['totals']['balance'] is None
    s.set_financial_evidence(r,True,True,'Reconciled all receipts')
    assert s.snapshot(r)['totals']['balance']==90000


def test_repair_search_finds_current_historical_customer_phone(tmp_path):
    s=ServiceStore(tmp_path/'project');c=s.save_client({'name':'Synthetic','phone':'+7 (000) 111-22-33'})
    e=s.save_equipment({'number':'E1','model':'Model','owner_id':c});r=s.create_repair(e)
    assert [row['id'] for row in s.repairs(query='111-22')]==[r]
    assert [row['id'] for row in s.repairs(query='70001112233')]==[r]

def test_unknown_line_prices_are_not_zero_in_documents_or_immutable_quotes(tmp_path):
    from docx import Document
    from docflow.service.documents import generate_document
    s,r,_=imported(tmp_path,{'final_total':'100'})
    snap=s.snapshot(r)
    assert snap['totals']['work'] is None
    assert snap['repair_items'][0]['price_known']==0
    q=freeze_estimate(s,r,'final','Synthetic operator')
    assert s.snapshot(r)['estimates'][0]['payload']['lines'][0]['price_known'] is False
    output=generate_document(s,r,'estimate_final',estimate_id=q)
    d=Document(s.root/output['path']) if isinstance(output,dict) else None
    # Inspect the actual registered document rather than assuming return type.
    registered=s.snapshot(r)['generated_documents'][0]
    d=Document(s.root/registered['path'])
    text=' '.join(p.text for p in d.paragraphs)+' '.join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert 'Неизвестно' in text and '100,00' in text


def test_v4_migration_keeps_a_rollback_copy_and_marks_legacy_finance_unknown(tmp_path):
    import sqlite3
    from docflow.service.database import MIGRATIONS
    root=tmp_path/'v4';root.mkdir();db=root/'service.db'
    with sqlite3.connect(db) as c:
        for v in range(1,5):
            for sql in MIGRATIONS[v]:c.execute(sql)
        c.execute("INSERT INTO clients(id,kind,name,created_at,updated_at) VALUES('c','organization','Synthetic','x','x')")
        c.execute("INSERT INTO equipment(id,number,model,owner_id,created_at,updated_at) VALUES('e','E','Model','c','x','x')")
        c.execute("INSERT INTO repair_orders(id,number,equipment_id,client_id,received_at,created_at,updated_at) VALUES('r','R','e','c','2020-01-01','x','x')")
        c.execute("INSERT INTO import_sources VALUES('s','source','sha','sheet','{}','x')")
        c.execute("INSERT INTO import_records VALUES('i','s','key','sha','r','{}','x')")
        c.execute('PRAGMA user_version=4')
    s=ServiceStore(root)
    assert s.snapshot('r')['totals']['paid'] is None
    old=list(root.glob('service.db.before-v5-*.bak'));assert len(old)==1
    with sqlite3.connect(old[0]) as c:
        assert c.execute('PRAGMA user_version').fetchone()[0]==4
        assert c.execute('SELECT number FROM repair_orders').fetchone()[0]=='R'
        assert not c.execute("SELECT 1 FROM sqlite_master WHERE name='repair_finance_evidence'").fetchall()

def test_unknown_prices_in_specification_and_contact_fields_in_custom_docx(tmp_path):
    from docx import Document
    from docflow.service.documents import generate_document
    from docflow.service.estimates import export_specification
    s,r,_=imported(tmp_path,{'final_total':'100','client_phone':'+7 (000) 111-22-33','client_contact':'Synthetic contact'})
    q=freeze_estimate(s,r,'final','Synthetic operator')
    csv_path=export_specification(s,q,tmp_path/'spec.csv')
    assert 'Неизвестно' in csv_path.read_text(encoding='utf-8-sig')
    template=tmp_path/'contact.docx';d=Document();d.add_paragraph('{{client_contact}} / {{client_phone}} / {{expenses_total}}');d.save(template)
    result=generate_document(s,r,'completion',template=template)
    assert 'Synthetic contact / +7 (000) 111-22-33' in Document(result).paragraphs[0].text
