import csv
import sqlite3
from datetime import date

import pytest
from docx import Document
from openpyxl import Workbook

from docflow.service.repository import ServiceStore
from docflow.service.history_import import preview_history, import_history, workbook_sheets
from docflow.service.docx_extract import export_review_csv
from docflow.service.validation import ServiceError

HEADERS=['Аппарат','Модель','Клиент','Ремонт','Дата','Жалоба','Дефект','Работы','Запчасти','Итог','Статус']
MAPPING=dict(zip(['equipment_number','model','client_name','repair_number','received_at','complaint','defects','works','parts','final_total','status'],HEADERS))
ROWS=[['900','Учебный насос','Учебный клиент А','101','2023-02-01','Нет давления','Учебный износ','Очистка','Учебное кольцо','100','Закрыт'],
      ['900','Учебный насос','Учебный клиент Б','305','2024-03-02','Повторная жалоба','Другой учебный износ','Настройка','Учебный фильтр','200','Закрыт'],
      ['900','Учебный насос','Учебный клиент А','900','2026-10-24','Вымышленная течь','Учебная неисправность','Разборка','Учебное кольцо','300','Закрыт']]


def write_csv(path,rows=ROWS,headers=HEADERS):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f,delimiter=';');w.writerow(headers);w.writerows(rows)
    return path


def options():
    return {'equipment_decisions':{'900':'new'},'client_decisions':{'учебный клиент а':'new','учебный клиент б':'new'},
            'current_owners':{'900':'учебный клиент б'},'status_map':{'Закрыт':'closed'},'default_status':'closed'}


def counts(s):
    with s.db.connection() as c:
        return {t:c.execute('SELECT COUNT(*) FROM '+t).fetchone()[0] for t in ('clients','equipment','repair_orders','repair_items','attachments','import_sources','import_records')}


def test_three_historical_repairs_fourth_repeat_owners_search_replay(tmp_path):
    s=ServiceStore(tmp_path/'project');p=write_csv(tmp_path/'history.csv');before=counts(s)
    plan=preview_history(s,p,MAPPING,**options())
    assert counts(s)==before and all(g['status']=='ready' for g in plan['groups'])
    report=import_history(s,plan,confirmed=True)
    assert report['imported']==3 and report['deferred']==0
    equipment=s.equipment()[0];history=s.repairs(equipment_id=equipment['id'],ascending=True)
    assert [r['number'] for r in history]==['101','305','900']
    assert [r['received_at'][:4] for r in history]==['2023','2024','2026']
    assert [r['client'] for r in history]==['Учебный клиент А','Учебный клиент Б','Учебный клиент А']
    assert s.get('clients',equipment['owner_id'])['name']=='Учебный клиент Б'
    fourth=s.create_repair(equipment['id'],'NEXT-REPAIR','Новый вымышленный ремонт')
    assert len(ServiceStore(s.root).repairs(equipment_id=equipment['id']))==4
    assert s.snapshot(fourth)['client']['name']=='Учебный клиент Б'
    assert len(s.repairs(query='учебное кольцо'))==2
    assert [r['number'] for r in s.repairs(date_from='2024-01-01',date_to='2024-12-31')]==['305']
    assert len(s.compare_repairs([history[0]['id'],history[1]['id']]))==2
    before=counts(s);again=import_history(s,plan,confirmed=True)
    assert again['status']=='no_changes' and again['already_imported']==3 and counts(s)==before


def test_no_auto_identity_and_explicit_deferral(tmp_path):
    s=ServiceStore(tmp_path/'project');p=write_csv(tmp_path/'history.csv')
    plan=preview_history(s,p,MAPPING,status_map={'Закрыт':'closed'},default_status='closed')
    assert all(g['status']=='needs_confirmation' for g in plan['groups'])
    with pytest.raises(ServiceError):import_history(s,plan,confirmed=True)
    assert counts(s)['repair_orders']==0
    bad=[list(r) for r in ROWS];bad[1][4]='not a date';p=write_csv(p,bad)
    plan=preview_history(s,p,MAPPING,**options())
    assert plan['groups'][1]['status']=='error'
    with pytest.raises(ServiceError):import_history(s,plan,confirmed=True)
    report=import_history(s,plan,confirmed=True,selected_keys=['101','900'])
    assert report['imported']==2 and report['deferred']==1
    assert counts(s)['repair_orders']==2


def test_rollback_with_real_sql_failure_and_copied_document(tmp_path):
    s=ServiceStore(tmp_path/'project');doc=tmp_path/'source.docx';d=Document();d.add_paragraph('Синтетический исходный акт');d.save(doc)
    rows=[r+['source.docx'] for r in ROWS];p=write_csv(tmp_path/'history.csv',rows,HEADERS+['Документ'])
    plan=preview_history(s,p,{**MAPPING,'source_document':'Документ'},**options(),document_root=tmp_path,copy_documents=True)
    assert all(g['status']=='ready' for g in plan['groups'])
    with s.db.connection(True) as c:
        c.execute("CREATE TRIGGER test_insert_error BEFORE INSERT ON repair_orders WHEN NEW.number='305' BEGIN SELECT RAISE(ABORT,'injected failure'); END")
    before=counts(s)
    with pytest.raises(ServiceError):import_history(s,plan,confirmed=True)
    assert counts(s)==before and not list((s.root/'attachments').rglob('*.*'))
    assert doc.is_file()


def test_file_change_and_source_conflict_never_overwrite_history(tmp_path):
    s=ServiceStore(tmp_path/'project');p=write_csv(tmp_path/'history.csv')
    plan=preview_history(s,p,MAPPING,**options(),source_label='Explicit stable source')
    changed=[list(r) for r in ROWS];changed[0][5]='Изменённая жалоба';write_csv(p,changed)
    with pytest.raises(ServiceError):import_history(s,plan,confirmed=True)
    assert counts(s)['clients']==0
    write_csv(p);plan=preview_history(s,p,MAPPING,**options(),source_label='Explicit stable source')
    import_history(s,plan,confirmed=True);before=counts(s)
    write_csv(p,changed);new=preview_history(s,p,MAPPING,**options(),source_label='Explicit stable source')
    assert new['groups'][0]['status']=='conflict'
    with pytest.raises(ServiceError):import_history(s,new,confirmed=True)
    assert counts(s)==before and s.repairs(query='Нет давления')


def test_explicit_existing_equipment_no_repair_number_as_identity(tmp_path):
    s=ServiceStore(tmp_path/'project');client=s.save_client({'name':'Текущий учебный владелец'})
    eq=s.save_equipment({'number':'FIXED-900','model':'Учебный насос','owner_id':client})
    p=write_csv(tmp_path/'history.csv',[ROWS[0]])
    with pytest.raises(ServiceError):preview_history(s,p,{**MAPPING,'equipment_number':'Ремонт'},**options())
    mapping={k:v for k,v in MAPPING.items() if k!='equipment_number'}
    plan=preview_history(s,p,mapping,equipment_decisions={'repair:101':eq},client_decisions={'учебный клиент а':'new'},status_map={'Закрыт':'closed'})
    assert plan['groups'][0]['status']=='ready'
    import_history(s,plan,confirmed=True)
    assert len(s.equipment())==1 and s.repairs()[0]['equipment_id']==eq
    assert s.get('equipment',eq)['owner_id']==client


def test_multiple_item_rows_and_two_similar_devices_are_not_merged(tmp_path):
    s=ServiceStore(tmp_path/'project')
    headers=['Аппарат','Модель','Клиент','Ремонт','Дата','Тип','Позиция','Кол-во','Цена','Итого']
    rows=[['900','Одинаковая модель','Учебный клиент','101','2023-01-01','work','Работа','1.5','1000.25','2200.58'],
          ['900','Одинаковая модель','Учебный клиент','101','2023-01-01','part','Деталь','2','200.10','2200.58'],
          ['900','Одинаковая модель','Учебный клиент','101','2023-01-01','delivery','Доставка','2','150','2200.58'],
          ['901','Одинаковая модель','Учебный клиент','102','2023-01-02','work','Другая работа','1','1','1']]
    p=write_csv(tmp_path/'items.csv',rows,headers)
    mapping=dict(zip(['equipment_number','model','client_name','repair_number','received_at','item_kind','item_name','quantity','price','final_total'],headers))
    plan=preview_history(s,p,mapping,equipment_decisions={'900':'new','901':'new'},client_decisions={'учебный клиент':'new'},default_status='closed')
    assert all(g['status']=='ready' for g in plan['groups'])
    import_history(s,plan,confirmed=True)
    assert len(s.equipment())==2 and len(s.repairs())==2
    snap=s.snapshot(s.repairs(query='101')[0]['id'])
    assert len(snap['repair_items'])==3 and snap['totals']['total']==220058 and snap['totals']['delivery']==30000


def test_xlsx_sheet_selection_dates_and_no_write_preview(tmp_path):
    s=ServiceStore(tmp_path/'project');p=tmp_path/'multiple sheets.xlsx';wb=Workbook();wb.active.title='Инструкция';wb.active.append(['Только инструкция'])
    ws=wb.create_sheet('История');ws.append(HEADERS)
    for row in ROWS:
        values=list(row);values[4]=date.fromisoformat(values[4]);ws.append(values)
    wb.save(p);wb.close()
    assert workbook_sheets(p)==['Инструкция','История']
    plan=preview_history(s,p,MAPPING,sheet='История',**options())
    assert len(plan['groups'])==3 and counts(s)['repair_orders']==0
    import_history(s,plan,confirmed=True);assert len(s.repairs())==3


def test_missing_owner_ambiguous_clients_and_unknown_status(tmp_path):
    s=ServiceStore(tmp_path/'project');p=write_csv(tmp_path/'history.csv')
    opt=options();opt.pop('current_owners')
    assert all(g['status']=='error' for g in preview_history(s,p,MAPPING,**opt)['groups'])
    c1=s.save_client({'name':'Учебный клиент А'});c2=s.save_client({'name':'Учебный клиент А'})
    plan=preview_history(s,p,MAPPING,equipment_decisions={'900':'new'},current_owners={'900':c1},status_map={'Закрыт':'closed'})
    assert plan['groups'][0]['client_candidates']==[c1,c2]
    assert plan['groups'][0]['status']=='needs_confirmation'
    opt=options();opt['status_map']={}
    assert all(g['status']=='error' for g in preview_history(s,p,MAPPING,**opt)['groups'])


def test_docx_controlled_extraction_keeps_anchors_and_does_not_import(tmp_path):
    source=tmp_path/'source.docx';d=Document();t=d.add_table(rows=3,cols=3);t.cell(0,0).merge(t.cell(0,1)).text='Проверяемый заголовок'
    t.cell(1,0).text='=HYPERLINK("bad")';t.cell(1,1).text='№101';t.cell(1,2).text='01.02.2023';t.cell(2,0).text='Учебный аппарат';d.save(source)
    output=tmp_path/'review.csv';proof=export_review_csv(source,output)
    assert not proof['database_writes'] and len(proof['records'])==3
    with output.open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f,delimiter=';'))
    assert rows[0]['source_fragment']=='word/document.xml/table[1]/row[1]'
    assert rows[0]['cell_01']=='Проверяемый заголовок' and rows[0]['cell_02']==''
    assert rows[1]['cell_01'].startswith("'=")
    assert not (tmp_path/'service.db').exists()


def test_import_document_must_stay_in_explicit_root(tmp_path):
    s=ServiceStore(tmp_path/'project');safe=tmp_path/'safe';safe.mkdir();outside=tmp_path/'outside.docx';d=Document();d.add_paragraph('Вымышленный акт');d.save(outside)
    p=write_csv(tmp_path/'history.csv',[ROWS[0]+['../outside.docx']],HEADERS+['Документ'])
    opt=options();opt['current_owners']={}
    plan=preview_history(s,p,{**MAPPING,'source_document':'Документ'},**opt,document_root=safe,copy_documents=True)
    assert plan['groups'][0]['status']=='error'
    with pytest.raises(ServiceError):import_history(s,plan,confirmed=True)
    assert counts(s)['repair_orders']==0
