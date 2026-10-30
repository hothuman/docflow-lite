"""Business text must be readable without mutating or discarding source evidence."""
import json
from types import SimpleNamespace

import pytest
from docx import Document

from docflow.service.record_text import readable_field, repair_display, preserve_field_edit
from docflow.service.repository import ServiceStore
from docflow.service.documents import generate_document
from docflow.service.backup import create_backup, restore_backup


@pytest.mark.parametrize('raw,expected', [
    ('Обычная жалоба', 'Обычная жалоба'),
    ('Строка 1\nСтрока 2', 'Строка 1\nСтрока 2'),
    ('Кириллица <не XML> & кавычки "так"', 'Кириллица <не XML> & кавычки "так"'),
    (1250, '1250'),
    (0, '0'),
    (json.dumps('Строка 1\nСтрока 2'), 'Строка 1\nСтрока 2'),
    (json.dumps('"Цитата оператора"'), '"Цитата оператора"'),
    (json.dumps(json.dumps('Строка 1\nСтрока 2')), 'Строка 1\nСтрока 2'),
    (r'C:\new\notes\двигатель.docx', r'C:\new\notes\двигатель.docx'),
    (r'Оператор намеренно написал \n и \\n', r'Оператор намеренно написал \n и \\n'),
    ('requires_identity_date_money_review написано оператором',
     'requires_identity_date_money_review написано оператором'),
])
def test_plain_text_and_complete_json_string_escape_rules(raw, expected):
    shown = readable_field(raw)
    assert shown.text == expected
    assert shown.raw == str(raw)
    assert not shown.warning


def source_envelope():
    return {
        'requires_identity_date_money_review': True,
        'source_key': 'synthetic-source-observation',
        'file_sha256': 'a' * 64,
        'mapped_raw': {'complaint': r'Не запускается\nПосле дождя',
                       'defects': r'Контакт повреждён\\nКоррозия',
                       'works': 'Очистка контакта', 'notes': 'Согласовать замену'},
        'cells': {'A2': {'value': 'raw value', 'type': 's'}},
    }


def test_source_envelope_projects_business_fields_and_never_mutates_input():
    envelope = source_envelope()
    raw = json.dumps(envelope, ensure_ascii=False)
    repair = {'number': '1000', 'complaint': '', 'defects': '', 'plan': '', 'notes': raw}
    original = dict(repair)
    shown = repair_display(repair)
    assert shown['complaint'] == 'Не запускается\nПосле дождя'
    assert shown['defects'] == 'Контакт повреждён\nКоррозия'
    assert shown['plan'] == ''
    assert shown['source_works'] == 'Очистка контакта'
    assert shown['notes'] == 'Согласовать замену'
    assert all('requires_identity_date_money_review' not in shown[k]
               for k in ('complaint', 'defects', 'plan', 'notes'))
    assert repair == original
    assert json.loads(raw) == envelope


def test_windows_paths_inside_source_wrappers_are_not_unescaped():
    raw = json.dumps({'source_key': 'test', 'notes': r'C:\new\notes\file.docx'}, ensure_ascii=False)
    assert readable_field(raw).text == r'C:\new\notes\file.docx'


@pytest.mark.parametrize('key',['works','work','repair','ремонт','работы'])
def test_performed_source_work_never_becomes_a_plan(key):
    raw=json.dumps({'source_fragment':'synthetic-row-1','mapped_raw':{key:'Заменили щётки'}},ensure_ascii=False)
    shown=repair_display({'notes':raw,'plan':''})
    assert shown['plan']==''
    assert shown['source_works']=='Заменили щётки'
    shown=repair_display({'notes':raw,'plan':'Проверить под нагрузкой'})
    assert shown['plan']=='Проверить под нагрузкой'
    assert shown['source_works']=='Заменили щётки'


def test_source_work_survives_note_edit_without_becoming_operational_work():
    raw=json.dumps(source_envelope(),ensure_ascii=False)
    edited=preserve_field_edit(raw,readable_field(raw).text,'Уточнено')
    shown=repair_display({'notes':edited,'plan':''})
    assert shown['source_works']=='Очистка контакта' and shown['plan']==''
    assert json.loads(edited)['original']==raw


def test_history_detail_labels_source_work_and_never_invents_plan():
    from docflow.ui.service.history import HistoryDialog
    raw=json.dumps(source_envelope(),ensure_ascii=False)
    row={'notes':raw,'received_at':'2026-10-29','complaint':'','defects':'','plan':'','item_text':''}
    captured=[]
    fake=SimpleNamespace(grid=SimpleNamespace(selected=lambda:row),detail=SimpleNamespace(setPlainText=captured.append))
    HistoryDialog.selected_detail(fake)
    assert 'Работы из источника: Очистка контакта' in captured[0]
    assert 'План: Очистка контакта' not in captured[0]


def test_json_business_object_keeps_numeric_zero_and_unknown_intentional_keys():
    assert readable_field('{"name":"Насос","amount":0,"custom":"Проверено"}').text == (
        'Название: Насос\nСумма: 0\ncustom: Проверено')
    record = {'notes': '{"complaint":"Не включается","notes":"Уточнить"}'}
    assert repair_display(record)['complaint'] == 'Не включается'
    assert repair_display(record)['notes'] == 'Уточнить'


@pytest.mark.parametrize('raw', ['{"notes":"Незакрытая строка', '[1,', '{не JSON}'])
def test_malformed_structure_is_safe_and_original_stays_available(raw):
    shown = readable_field(raw)
    assert shown.warning and 'Первоисточник' in shown.text
    assert shown.raw == raw
    assert raw not in shown.text
    edited = preserve_field_edit(raw, shown.text, 'Проверено оператором')
    assert json.loads(edited)['original'] == raw
    assert readable_field(edited).text == 'Проверено оператором'


def test_unchanged_projection_save_retains_exact_original_and_explicit_edit_keeps_evidence():
    raw = json.dumps(source_envelope(), ensure_ascii=False, indent=2)
    display = readable_field(raw).text
    assert preserve_field_edit(raw, display, display) == raw
    edited = preserve_field_edit(raw, display, 'Добавлена деловая заметка')
    content = json.loads(edited)
    assert content['original'] == raw
    assert json.loads(content['original']) == source_envelope()
    assert readable_field(edited).text == 'Добавлена деловая заметка'
    assert repair_display({'notes':edited})['complaint']=='Не запускается\nПосле дождя'
    assert preserve_field_edit('', 'Из исходных данных', 'Из исходных данных', 'complaint') == ''


def test_deep_json_is_bounded_and_does_not_crash_gui_projection():
    assert readable_field('[' * 2000 + '0' + ']' * 2000).warning
    nested = {'value': {'value': {'value': {'value': {'value': {'value': {'value': 'deep'}}}}}}}
    assert 'первоисточнике' in readable_field(json.dumps(nested)).text


def test_generated_docx_is_business_text_and_backup_preserves_exact_provenance(tmp_path):
    store = ServiceStore(tmp_path/'project')
    client = store.save_client({'name': 'Синтетический клиент'})
    eq = store.save_equipment({'number': 'APP-100', 'model': 'Учебная модель', 'owner_id': client})
    repair = store.create_repair(eq, '1000')
    raw = json.dumps(source_envelope(), ensure_ascii=False, indent=2)
    store.save_repair(repair, {'notes': raw})
    template = tmp_path/'fields.docx'
    doc = Document()
    for key in ('complaint', 'defects', 'plan', 'notes','source_works'):
        doc.add_paragraph('{{'+key+'}}')
    doc.save(template)
    generated = generate_document(store, repair, 'intake', template)
    output = '\n'.join(p.text for p in Document(generated).paragraphs)
    assert 'Не запускается\nПосле дождя' in output
    assert 'Контакт повреждён\nКоррозия' in output
    assert 'Очистка контакта' in output and 'Согласовать замену' in output
    assert 'requires_identity_date_money_review' not in output
    assert 'mapped_raw' not in output and 'source_key' not in output
    assert store.snapshot(repair)['repair']['notes'] == raw
    backup = create_backup(store, tmp_path/'snapshot.zip')
    restored = restore_backup(backup, tmp_path/'restored')
    assert ServiceStore(restored).snapshot(repair)['repair']['notes'] == raw


def test_builtin_combined_and_history_include_labelled_source_work_without_changing_data(tmp_path):
    from docflow.service.documents import _history
    store=ServiceStore(tmp_path/'project')
    client=store.save_client({'name':'Синтетический клиент'})
    equipment=store.save_equipment({'number':'DEMO-SOURCE-WORK','model':'Учебная модель','owner_id':client})
    first=store.create_repair(equipment,'2001')
    raw=json.dumps(source_envelope(),ensure_ascii=False)
    store.save_repair(first,{'notes':raw})
    output=generate_document(store,first,'combined')
    text='\n'.join(p.text for p in Document(output).paragraphs)
    assert 'Работы из источника\nОчистка контакта' in text
    assert 'План: Очистка контакта' not in text
    assert 'requires_identity_date_money_review' not in text
    second=store.create_repair(equipment,'2002')
    history=_history(store,store.snapshot(second)['repair'])
    assert history[0]['history_works']=='Работы из источника: Очистка контакта'
    assert store.snapshot(first)['repair']['plan']==''
    assert store.snapshot(first)['repair']['notes']==raw
    ordinary=generate_document(store,second,'combined')
    assert 'Работы из источника' not in '\n'.join(p.text for p in Document(ordinary).paragraphs)
