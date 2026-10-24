"""CP3: owner-observed numbers, immutable historical identity and replay."""
import pytest
from docx import Document
from docflow.service.repository import ServiceStore
from docflow.service.estimates import save_draft,freeze_estimate,record_event,get_estimate,export_specification
from docflow.service.documents import generate_document
from docflow.service.validation import ServiceError

def project(tmp_path):
 s=ServiceStore(tmp_path/'Синтетический проект')
 c=s.save_client({'name':'Учебная мастерская','kind':'organization'})
 e=s.save_equipment({'number':'900','model':'Учебный аппарат','owner_id':c})
 return s,e,s.create_repair(e,'DEMO-ACCEPTANCE')

def test_owner_observed_0_7500_9670_2170_scenario(tmp_path):
 s,e,r=project(tmp_path)
 v1=freeze_estimate(s,r,'preliminary','Тестовый оператор')
 assert s.snapshot(r)['estimates'][0]['total']==0 and s.snapshot(r)['totals']['balance']==0
 lines=[{'kind':kind,'name':name,'quantity':'1','price':price} for kind,name,price in [('work','Учебная работа','4000'),('part','Учебная деталь','3000'),('delivery','Учебная доставка','500')]]
 save_draft(s,r,lines,expected_revision=0)
 v2=freeze_estimate(s,r,'preliminary','Тестовый оператор','Уточнён учебный расчёт')
 assert s.snapshot(r)['estimate_comparison']['preliminary_total']==750000
 assert s.snapshot(r)['totals']['balance']==0
 record_event(s,v2,'approved','Тестовый оператор','Синтетическая ручная отметка, не сообщение клиенту')
 for kind,name,price in [('work','Фактическая учебная работа','5000'),('part','Фактическая учебная деталь','4000'),('delivery','Учебная доставка','500'),('other','Учебный расход','170')]:
  s.save_item(r,{'kind':kind,'name':name,'price':price,'change_reason':'Уточнён фактический объём'})
 snap=s.snapshot(r);assert snap['totals']['total']==967000
 assert snap['estimate_comparison']['difference']==217000 and snap['totals']['balance']==967000
 final=freeze_estimate(s,r,'final','Тестовый оператор','Фактический объём после дефектовки')
 s.add_payment(r,'10000',method='Учебная запись');assert s.snapshot(r)['totals']['overpaid']==33000
 s.add_payment(r,'330','refund',method='Учебная запись');assert s.snapshot(r)['totals']['balance']==0
 for kind,key,total in [('estimate_preliminary',v2,'7 500,00 ₽'),('estimate_final',final,'9 670,00 ₽')]:
  doc=Document(generate_document(s,r,kind,estimate_id=key))
  text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
  assert total in text
  assert export_specification(s,key,tmp_path/(kind+'.csv')).is_file()
 with s.db.connection() as c:
  assert get_estimate(c,v1)['total']==0 and get_estimate(c,v2)['total']==750000
  assert get_estimate(c,final)['total']==967000
 again=ServiceStore(s.root);assert again.snapshot(r)['totals']['paid']==967000

@pytest.mark.parametrize('changed',[{'method':'Другой способ'},{'comment':'Другое основание'},{'occurred_at':'2026-10-25'}])
def test_payment_replay_rejects_changed_metadata_without_mutation(tmp_path,changed):
 s,e,r=project(tmp_path)
 original={'method':'Учебный способ','comment':'Учебное основание','occurred_at':'2026-10-24'}
 key=s.add_payment(r,'100',request_id='same-operation',**original)
 before=s.snapshot(r)
 with pytest.raises(ServiceError,match='другой записи'):
  s.add_payment(r,'100',request_id='same-operation',**{**original,**changed})
 assert s.snapshot(r)==before and s.add_payment(r,'100',request_id='same-operation',**original)==key
