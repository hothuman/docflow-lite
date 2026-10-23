import io,json,sqlite3,zipfile,threading
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
import pytest
from docx import Document
from docflow.service.repository import ServiceStore
from docflow.service.database import MIGRATIONS,Database,SCHEMA_VERSION
from docflow.service.validation import ServiceError,money,quantity,rub
from docflow.service.demo import seed_demo
from docflow.service.documents import generate_document,render_service,create_templates,KINDS
from docflow.service.attachments import add_attachment,attachment_path,remove_attachment,inspect_file,safe_path
from docflow.service.backup import create_backup,restore_backup
from docflow.service.importers import preview_equipment,import_equipment

@pytest.fixture
def store(tmp_path):
 s=ServiceStore(tmp_path/'Кириллица с пробелами');ids=seed_demo(s.root);return s,ids

def test_repeat_repair_identity_owner_history_persist(store):
 s,ids=store;c2=s.save_client({'name':'Новый вымышленный владелец'})
 e=s.get('equipment',ids['equipment']);s.save_equipment({**e,'owner_id':c2},e['id'])
 third=s.create_repair(e['id'],client_id=c2)
 again=ServiceStore(s.root);rows=again.repairs(equipment_id=e['id'])
 assert len(rows)==3 and rows[0]['id']==third
 assert again.snapshot(ids['repair'])['client']['id']==ids['client']
 assert again.snapshot(third)['client']['id']==c2
 assert again.dashboard()['equipment']==1

def test_numbers_and_atomic_duplicate_import(store):
 s,ids=store
 with pytest.raises(ServiceError):s.save_equipment({'number':'demo-001','model':'x','owner_id':ids['client']})
 with pytest.raises(ServiceError):s.create_repair(ids['equipment'],'demo-r-001')
 with pytest.raises(ServiceError):import_equipment(s,[{'number':'NEW','model':'A'},{'number':'DEMO-001','model':'B'}],ids['client'])
 assert len(s.equipment())==1

def test_idempotent_double_click_and_conflict(store):
 s,ids=store
 def create():return s.create_repair(ids['equipment'],request_id='double-click-key')
 with ThreadPoolExecutor(2) as pool:assert len(set(pool.map(lambda _:create(),range(2))))==1
 r=s.snapshot(ids['repair'])['repair'];s.save_repair(r['id'],{'defects':'Правка'},r['revision'])
 with pytest.raises(ServiceError):s.save_repair(r['id'],{'defects':'Потерянная правка'},r['revision'])
 assert s.snapshot(r['id'])['repair']['defects']=='Правка'

def test_status_and_finance_are_independent(store):
 s,ids=store;r=ids['repair'];s.save_repair(r,{'status':'ready'})
 v=s.snapshot(r);assert v['totals']['balance']==170050 and v['repair']['status']=='ready'
 s.save_repair(r,{'status':'issued','issued_at':'2026-10-23','handover':'Демонстрационная отметка, не подпись'})
 s.save_repair(r,{'status':'diagnostics','status_note':'Исправление ошибочного статуса'})
 h=s.snapshot(r)['repair_status_history'];assert len(h)==5 and h[-1]['note']=='Исправление ошибочного статуса'
 assert s.snapshot(r)['totals']['paid']==100000

@pytest.mark.parametrize('value',['-1','NaN','Infinity','1.001','1e99','abc'])
def test_invalid_money(value):
 with pytest.raises(ServiceError):money(value)

@pytest.mark.parametrize('value',['0','-1','NaN','1.12345','1000000001'])
def test_invalid_quantity(value):
 with pytest.raises(ServiceError):quantity(value)

def test_decimal_rounding_refund_and_void(store):
 s,ids=store;r=s.create_repair(ids['equipment'])
 item=s.save_item(r,{'kind':'work','name':'Точность','quantity':'0.5','price':'0.01'})
 assert s.snapshot(r)['totals']['total']==1
 pay=s.add_payment(r,'1.00',request_id='payment-click')
 assert s.add_payment(r,'1.00',request_id='payment-click')==pay
 with pytest.raises(ServiceError):s.add_payment(r,'2',request_id='payment-click')
 assert s.snapshot(r)['totals']['overpaid']==99
 refund=s.add_payment(r,'0.50','refund');assert s.snapshot(r)['totals']['paid']==50
 with pytest.raises(ServiceError):s.void_payment(pay,'Ошибочная оплата')
 assert not s.get('payments',pay)['voided']
 s.void_payment(refund,'Ошибочный возврат');s.void_payment(pay,'Исправление');assert s.snapshot(r)['totals']['paid']==0
 assert len(s.snapshot(r)['payments'])==2
 s.save_repair(r,{'discount':'0.01'})
 with pytest.raises(ServiceError):s.remove_item(item)
 s.save_repair(r,{'discount':'0'});s.remove_item(item)
 assert s.snapshot(r)['totals']['total']==0 and rub(-105)=='−1,05 ₽'

def test_archive_preserves_history(store):
 s,ids=store
 with pytest.raises(ServiceError):s.archive('equipment',ids['equipment'])
 s.save_repair(ids['repair'],{'status':'closed'});s.archive('equipment',ids['equipment'])
 assert not s.equipment() and len(s.equipment(archived=True))==1 and len(s.repairs())==2
 s.archive('equipment',ids['equipment'],False);assert s.equipment()

def test_migrate_v1_preserves_rows_and_backup(tmp_path):
 path=tmp_path/'old.db'
 with sqlite3.connect(path) as c:
  for q in MIGRATIONS[1]:c.execute(q)
  c.execute("INSERT INTO clients(id,kind,name,created_at,updated_at) VALUES('c','organization','Legacy synthetic','x','x')")
  c.execute('PRAGMA user_version=1')
 db=Database(path)
 with db.connection() as c:assert c.execute('PRAGMA user_version').fetchone()[0]==SCHEMA_VERSION;assert c.execute('SELECT name FROM clients').fetchone()[0]=='Legacy synthetic'
 assert list(tmp_path.glob('old.db.before-*.bak'))

def test_failed_migration_rolls_back(tmp_path,monkeypatch):
 path=tmp_path/'old.db'
 with sqlite3.connect(path) as c:
  for q in MIGRATIONS[1]:c.execute(q)
  c.execute('PRAGMA user_version=1')
 monkeypatch.setitem(MIGRATIONS,2,['CREATE TABLE transient(x)','BROKEN SQL'])
 with pytest.raises(ServiceError):Database(path)
 with sqlite3.connect(path) as c:
  assert c.execute('PRAGMA user_version').fetchone()[0]==1
  assert not c.execute("SELECT 1 FROM sqlite_master WHERE name='transient'").fetchone()

def test_corrupted_and_future_database(tmp_path):
 p=tmp_path/'broken.db';p.write_bytes(b'not sqlite')
 with pytest.raises(ServiceError):Database(p)
 p.unlink()
 with sqlite3.connect(p) as c:c.execute('PRAGMA user_version=999')
 with pytest.raises(ServiceError):Database(p)

@pytest.mark.parametrize('kind',list(KINDS))
def test_documents_versions_and_no_tokens(store,kind):
 s,ids=store
 p=generate_document(s,ids['repair'],kind);p2=generate_document(s,ids['repair'],kind)
 assert p!=p2 and p.read_bytes()!=b''
 d=Document(p);txt='\n'.join([x.text for x in d.paragraphs]+[cell.text for t in d.tables for row in t.rows for cell in row.cells])
 assert '{{' not in txt and '[[items:' not in txt and 'DEMO-R-002' in txt
 if kind in ('combined','completion'):assert any(len(t.rows)==2 for t in d.tables);assert 'Настройка узла' in txt
 assert len([x for x in s.snapshot(ids['repair'])['generated_documents'] if x['kind']==kind])==2

def test_empty_many_rows_and_custom_split_runs(store,tmp_path):
 s,ids=store;r=s.create_repair(ids['equipment']);path=generate_document(s,r,'completion')
 assert [len(t.rows) for t in Document(path).tables]==[1,1]
 for i in range(65):s.save_item(r,{'kind':'work','name':f'Вымышленная работа {i}','description':'Длинное описание '*12,'quantity':'1.125','price':'100.25'})
 path=generate_document(s,r,'completion');assert len(Document(path).tables[0].rows)==66
 d=Document();d.add_paragraph('{{repair_number}}');d.sections[0].header.paragraphs[0].text='Сервис {{client_name}}'
 t=d.add_table(rows=2,cols=2);t.style='Table Grid';t.cell(0,0).text='Работа'
 p=t.cell(1,0).paragraphs[0];p.add_run('[[items:');p.add_run('work]]{{item_');p.add_run('name}}');t.cell(1,1).text='{{item_total}}'
 template=tmp_path/'custom.docx';d.save(template)
 out=generate_document(s,r,'completion',template);result=Document(out)
 assert len(result.tables[0].rows)==66 and result.tables[0].cell(1,0).text=='Вымышленная работа 0'
 assert '{{' not in result.sections[0].header.paragraphs[0].text

def test_unsafe_docx_rejected(store,tmp_path):
 s,ids=store;templates=create_templates(tmp_path/'forms');source=templates['intake']
 path=tmp_path/'bad.docx'
 with zipfile.ZipFile(source) as a,zipfile.ZipFile(path,'w') as z:
  for n in a.namelist():z.writestr(n,a.read(n))
  z.writestr('word/vbaProject.bin',b'danger')
 with pytest.raises(Exception):generate_document(s,ids['repair'],'intake',path)
 assert not s.snapshot(ids['repair'])['generated_documents']

@pytest.mark.parametrize('name',['../outside','/outside','attachments/../../outside','attachments\\outside','C:/file'])
def test_path_traversal(tmp_path,name):
 with pytest.raises(ServiceError):safe_path(tmp_path,name)

def test_real_image_video_missing_and_backup_restore(store,tmp_path):
 from PySide6.QtGui import QImage,QColor
 s,ids=store;r=ids['repair'];photo=tmp_path/'Фото с пробелами.png';im=QImage(200,150,QImage.Format.Format_RGB32);im.fill(QColor('#873047'));assert im.save(str(photo))
 key=add_attachment(s,r,photo,'Вымышленная схема');assert attachment_path(s,key).is_file()
 video=Path(__file__).resolve().parents[2]/'site/dist/assets/demo.mp4';v=add_attachment(s,r,video)
 generated=generate_document(s,r,'combined',image_ids=[key]);assert len(Document(generated).inline_shapes)==1
 b=create_backup(s,tmp_path/'snapshot.zip');restore_backup(b,tmp_path/'Восстановленный проект')
 again=ServiceStore(tmp_path/'Восстановленный проект');assert len(again.snapshot(r)['attachments'])==2
 assert attachment_path(again,v).read_bytes()==video.read_bytes()
 attachment_path(s,key).write_bytes(b'changed')
 with pytest.raises(ServiceError):attachment_path(s,key)
 with pytest.raises(ServiceError):create_backup(s,tmp_path/'incomplete.zip')
 remove_attachment(s,key);assert len(s.snapshot(r)['attachments'])==1
 with pytest.raises(ServiceError):restore_backup(b,again.root)

def test_media_extension_and_corruption(tmp_path):
 for name,raw in [('x.exe',b'MZ\0\0'),('x.png',b'not image'),('x.mp4',b'\0\0\0\x20ftyp')]:
  p=tmp_path/name;p.write_bytes(raw)
  with pytest.raises(ServiceError):inspect_file(p)

def test_backup_traversal_corruption_and_destination_preserved(store,tmp_path):
 s,ids=store;backup=create_backup(s,tmp_path/'b.zip');original=backup.read_bytes()
 with pytest.raises(ServiceError):create_backup(s,backup)
 assert backup.read_bytes()==original
 corrupt=tmp_path/'corrupt.zip';corrupt.write_bytes(b'not zip')
 with pytest.raises(ServiceError):restore_backup(corrupt,tmp_path/'new')
 assert not (tmp_path/'new').exists()
 malicious=tmp_path/'evil.zip'
 with zipfile.ZipFile(malicious,'w') as z:
  z.writestr('../escape',b'x');z.writestr('manifest.json',json.dumps({'format':'docflow-service-backup-1','schema':2,'files':{'../escape':{'sha256':'x','size':1}}}))
 with pytest.raises(ServiceError):restore_backup(malicious,tmp_path/'evil-dest')
 assert not (tmp_path/'escape').exists() and not (tmp_path/'evil-dest').exists()

def test_import_preview_explicit_and_atomic(store,tmp_path):
 s,ids=store;f=tmp_path/'equipment.csv';f.write_text('Номер,Модель\nNEW-1,Тест А\nNEW-2,Тест Б\n',encoding='utf-8-sig')
 rows=preview_equipment(f,{'number':'Номер','model':'Модель'});assert len(rows)==2
 import_equipment(s,rows,ids['client']);assert len(s.equipment())==3
 with pytest.raises(ServiceError):import_equipment(s,rows,ids['client'])
 assert len(s.equipment())==3

def test_unicode_identifier_uniqueness(store):
 s,ids=store;s.save_equipment({'number':'АППАРАТ-А','model':'Тест','owner_id':ids['client']})
 with pytest.raises(ServiceError):s.save_equipment({'number':'аппарат-а','model':'Другой','owner_id':ids['client']})
 s.create_repair(ids['equipment'],'РЕМОНТ-А')
 with pytest.raises(ServiceError):s.create_repair(ids['equipment'],'ремонт-а')

def test_backup_disk_full_keeps_source_and_removes_partial(store,tmp_path,monkeypatch):
 import docflow.service.backup as b
 s,ids=store;before=s.snapshot(ids['repair'])
 original=b.shutil.copyfileobj
 def full(src,dst,*args):
  if args:return original(src,dst,*args)
  dst.write(b'partial');raise OSError('No space left on device')
 monkeypatch.setattr(b.shutil,'copyfileobj',full)
 with pytest.raises(OSError):b.create_backup(s,tmp_path/'full.zip')
 assert not (tmp_path/'full.zip').exists() and s.snapshot(ids['repair'])==before
