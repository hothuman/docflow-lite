from pathlib import Path
import hashlib
import pytest
from docx import Document
from docflow.service.repository import ServiceStore
from docflow.service.documents import generate_document
from docflow.service.profiles import save_profile,NEUTRAL
from docflow.service.estimates import save_draft,freeze_estimate,get_estimate
from docflow.service.validation import ServiceError

def setup(tmp_path):
    store=ServiceStore(tmp_path/'project');repair=store.intake(client={'name':'Учебный'},equipment={'number':'A-01','model':'Учебная мойка'});return store,repair

def test_neutral_form_never_silently_outputs_empty_executor(tmp_path):
    store,repair=setup(tmp_path);path=generate_document(store,repair,'intake');doc=Document(path)
    assert 'исполнитель не указан' in '\n'.join(p.text for p in doc.paragraphs).casefold()
    assert doc.styles['Normal'].font.size.pt>=11

def test_profile_assignment_does_not_rewrite_old_quotes_or_documents(tmp_path):
    store,repair=setup(tmp_path);save_draft(store,repair,[{'kind':'work','name':'Учебная работа','price':'100'}]);estimate=freeze_estimate(store,repair,'preliminary','Учебный')
    old=generate_document(store,repair,'intake');old_sha=hashlib.sha256(old.read_bytes()).hexdigest()
    with store.db.connection() as c:before=get_estimate(c,estimate)
    revision=store.snapshot(repair)['repair']['revision'];key=save_profile(store,{**NEUTRAL,'name':'Учебный профиль','organization_name':'Вымышленная мастерская'},repair_id=repair,expected_revision=revision,activate=True)
    assert store.snapshot(repair)['repair']['profile_id']==key
    assert hashlib.sha256(old.read_bytes()).hexdigest()==old_sha
    with store.db.connection() as c:assert get_estimate(c,estimate)==before
    new=generate_document(store,repair,'intake');assert 'Вымышленная мастерская' in '\n'.join(p.text for p in Document(new).paragraphs)

def test_profile_assignment_conflict_rolls_back_new_profile(tmp_path):
    store,repair=setup(tmp_path)
    with pytest.raises(ServiceError):save_profile(store,{**NEUTRAL,'name':'Учебный','organization_name':'Вымышленная мастерская'},repair_id=repair,expected_revision=999)
    with store.db.connection() as c:assert c.execute('SELECT COUNT(*) FROM organization_profiles').fetchone()[0]==0

def test_document_preflight_uses_frozen_quote_profile_not_current_profile(tmp_path):
    from docflow.service.document_readiness import assess
    store,repair=setup(tmp_path);save_draft(store,repair,[]);freeze_estimate(store,repair,'preliminary','Учебный')
    revision=store.snapshot(repair)['repair']['revision'];save_profile(store,{**NEUTRAL,'name':'Учебный','organization_name':'Вымышленная мастерская','contact':'Учебный контакт','details':'Учебные реквизиты'},repair_id=repair,expected_revision=revision)
    snapshot=store.snapshot(repair);assert not assess(snapshot,'intake')['missing']
    quote=assess(snapshot,'estimate_preliminary');assert quote['quote_snapshot'] and 'название исполнителя' in quote['missing']
    assert assess(snapshot,'intake','custom.docx')['blocked'] is False

def test_wizard_cancel_and_explicit_assignment(app,tmp_path):
    from test_stage1_upgrade_copy import idle,dispose
    from docflow.ui.service.workspace import ServiceWorkspace
    from docflow.ui.service.repair_editor import RepairDialog
    from docflow.ui.service.organization_setup import OrganizationWizard
    w=ServiceWorkspace(root=tmp_path/'p');idle(app,w);repair=w.store.intake(client={'name':'Учебный'},equipment={'number':'A-01','model':'Учебная'})
    try:
        editor=RepairDialog(w,w.store.snapshot(repair));editor.show();wizard=OrganizationWizard(w,editor);wizard.show();wizard.reject();assert not w.store.snapshot(repair)['repair']['profile_id']
        from PySide6.QtWidgets import QWizard
        wizard=OrganizationWizard(w,editor);wizard.show();assert not wizard.button(QWizard.WizardButton.NextButton).isEnabled()
        wizard.organization.setText('Вымышленная мастерская');wizard.contact.setText('Учебный контакт');wizard.details.setPlainText('Учебные реквизиты');assert wizard.button(QWizard.WizardButton.NextButton).isEnabled();wizard.button(QWizard.WizardButton.NextButton).click();assert wizard.currentId()==1;wizard.next();wizard.accept();idle(app,w)
        assert wizard.saved_profile_id and editor.snapshot['profile']['organization_name']=='Вымышленная мастерская'
        assert 'не указан' not in editor.profile_banner.text();editor.accept()
    finally:dispose(app,w)

def test_incomplete_profile_blocks_custom_form_without_writing(app,tmp_path):
    from test_stage1_upgrade_copy import idle,dispose
    from docflow.ui.service.workspace import ServiceWorkspace
    from docflow.ui.service.repair_editor import RepairDialog
    w=ServiceWorkspace(root=tmp_path/'p');idle(app,w);repair=w.store.intake(client={'name':'Учебный'},equipment={'number':'A-01','model':'Учебная'})
    try:
        editor=RepairDialog(w,w.store.snapshot(repair));editor.template_path.setText(str(tmp_path/'custom.docx'));editor.generate();assert not w.store.documents() and 'Документ не создан' in editor.document_status.text();editor.accept()
    finally:dispose(app,w)

from test_stage1_upgrade_copy import app
