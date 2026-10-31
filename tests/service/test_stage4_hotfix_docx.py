import hashlib
from docx import Document
from docx.oxml.ns import qn
from docflow.service.repository import ServiceStore
from docflow.service.documents import generate_document,create_templates

def test_neutral_signing_block_keeps_financial_context(tmp_path):
    files=create_templates(tmp_path,modern=True);d=Document(files['completion'])
    for prefix in ('Внутренние записи','Производственный статус:','Проект формы.'):
        p=next(p for p in d.paragraphs if p.text.startswith(prefix))
        assert p.paragraph_format.keep_with_next is True,prefix
    assert next(p for p in d.paragraphs if p.text.startswith('Передал ')).paragraph_format.keep_with_next is not True

def test_long_rows_old_versions_and_custom_template_preserved(tmp_path):
    s=ServiceStore(tmp_path/'p');r=s.intake(client={'name':'QA'},equipment={'number':'QA','model':'QA'})
    first=generate_document(s,r,'completion');sha=hashlib.sha256(first.read_bytes()).hexdigest()
    for i in range(48):s.save_item(r,{'kind':'work','name':f'QA {i:02d}','description':'QA длинное описание. '*12,'quantity':'1','price':'1.01'})
    second=generate_document(s,r,'completion');d=Document(second);content='\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
    for i in range(48):assert f'QA {i:02d}' in content
    assert hashlib.sha256(first.read_bytes()).hexdigest()==sha and first!=second
    for row in d.tables[0].rows[1:]:
        assert row._tr.trPr.find(qn('w:cantSplit')) is not None
        assert all(p.paragraph_format.keep_together for c in row.cells for p in c.paragraphs)
    custom=tmp_path/'custom.docx';x=Document();x.add_paragraph('{{repair_number}}');x.save(custom);old=custom.read_bytes();generated=generate_document(s,r,'completion',template=custom)
    assert custom.read_bytes()==old and Document(generated).paragraphs[0].text==s.snapshot(r)['repair']['number']
