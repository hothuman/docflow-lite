"""Black-box regression cases added by the independent verification pass."""
import io
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import time
import sys
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook

from docflow.data import load_table
from docflow.errors import DocFlowError
from docflow.profile import Profile


def sheet_with_dimension(tmp_path, dimension, *, rows=3, far_cell=None):
    p = tmp_path / "dimensions.xlsx"
    wb = Workbook()
    wb.active.append(["name", "email"])
    for i in range(rows):
        wb.active.append([f"Client {i}", f"{i}@example.com"])
    if far_cell:
        wb.active[far_cell] = "extra"
    wb.save(p)
    out = io.BytesIO()
    with zipfile.ZipFile(p) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for entry in src.infolist():
            data = src.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"\s*/>', f'<dimension ref="{dimension}"/>'.encode(), data)
            dst.writestr(entry, data)
    p.write_bytes(out.getvalue())
    return p


@pytest.mark.parametrize("dimension", ["A1:B2", "A1:A4", "A1:XFD1048576"])
def test_real_cells_survive_incorrect_dimension_metadata(tmp_path, dimension):
    table = load_table(sheet_with_dimension(tmp_path, dimension))
    assert table.headers == ("name", "email")
    assert len(table.records) == 3
    assert table.records[-1].values["email"] == "2@example.com"


def test_hidden_large_row_not_silently_ignored(tmp_path):
    p = sheet_with_dimension(tmp_path, "A1:B2", far_cell="A1000000")
    with pytest.raises(DocFlowError):
        load_table(p)


def test_hidden_large_column_not_silently_ignored(tmp_path):
    p = sheet_with_dimension(tmp_path, "A1:B2", far_cell="CW2")
    with pytest.raises(DocFlowError):
        load_table(p)


def test_repo_does_not_ignore_website_source(tmp_path):
    if not shutil.which("git"):
        pytest.skip("Git unavailable")
    source = Path(__file__).resolve().parents[1]
    shutil.copy(source / ".gitignore", tmp_path / ".gitignore")
    (tmp_path / "site/dist").mkdir(parents=True)
    (tmp_path / "site/dist/index.html").write_text("test")
    subprocess.run(["git", "init", "-q", str(tmp_path)], check=True)
    result = subprocess.run(["git", "check-ignore", "site/dist/index.html"], cwd=tmp_path, capture_output=True)
    assert result.returncode == 1, "Website would be silently excluded by git add"


def test_gui_generation_uses_the_record_that_was_previewed(tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication
    from docflow.gui import MainWindow
    from docflow.release_check import wait_idle
    app = QApplication.instance() or QApplication([])
    source = tmp_path / "data.csv"
    source.write_text("name\nOriginal Client\n", encoding="utf-8")
    template = tmp_path / "template.docx"
    d = Document(); d.add_paragraph("{{name}}"); d.save(template)
    window = MainWindow()
    window.apply_profile(Profile(source, [template], {"name":"name"}))
    wait_idle(window)
    assert "Original Client" in window.preview_text.toPlainText()
    # An external process changes both inputs after the user sees a preview.
    source.write_text("name\nUnexpected Client\n", encoding="utf-8")
    d = Document(); d.add_paragraph("MODIFIED {{name}}"); d.save(template)
    window.output_path.setText(str(tmp_path / "out"))
    window.start_generation()
    deadline = time.monotonic() + 20
    while (window.worker.isRunning() or window.last_result is None) and time.monotonic() < deadline:
        app.processEvents(); time.sleep(.01)
    window.worker.wait(1000); app.processEvents()
    try:
        assert window.last_result is not None
        output = next(window.last_result.directory.rglob("*.docx"))
        assert Document(output).paragraphs[0].text == "Original Client"
    finally:
        window.close()


def test_release_can_be_repacked_with_verifiable_hashes(tmp_path):
    source = Path(__file__).resolve().parents[1]
    if not (source / 'scripts/package_release.py').is_file() or not (source / 'site/dist/index.html').is_file():
        pytest.skip('Release packaging integration requires the full developer checkout with website source')
    checkout = tmp_path / "source"
    declared=json.loads((source/'FILE_HASHES.json').read_text(encoding='utf-8'))
    for name in [*declared,'FILE_HASHES.json']:
        target=checkout/name;target.parent.mkdir(parents=True,exist_ok=True)
        shutil.copy2(source/name,target)
    # Untracked private/stale inputs must never enter a source/starter or public demo.
    private_marker=b'UNTRACKED_CUSTOMER_'+bytes(range(20))
    for name in ('.env','client-work/customer.txt','release-old/stale.zip'):
        target=checkout/name;target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(private_marker)
    result = subprocess.run([sys.executable, 'scripts/package_release.py'], cwd=checkout, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    archive = next((tmp_path / 'deliverables').glob('*Starter.zip'))
    with zipfile.ZipFile(archive) as z:
        assert len(z.namelist()) == len(set(z.namelist()))
        assert not any(private_marker in z.read(n) for n in z.namelist())
        checksums = json.loads(z.read('docflow-lite/FILE_HASHES.json'))
        assert all(hashlib.sha256(z.read('docflow-lite/' + path)).hexdigest() == digest for path, digest in checksums.items())
