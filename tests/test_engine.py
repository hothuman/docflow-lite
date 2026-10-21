import io
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from docflow.batch import generate, prepare
from docflow.cli import demo_profile, main
from docflow.data import Record, load_table
from docflow.errors import DocFlowError
from docflow.profile import Profile
from docflow.safety import check_archive, csv_safe, safe_name
from docflow.templates import Template, W, parse
from docflow.validation import Rule, values_for


def template(tmp_path, text="{{name}}", filename="template.docx"):
    path = tmp_path / filename
    d = Document()
    d.add_paragraph(text)
    d.save(path)
    return path


def rewrite(path, transform):
    out = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(info, transform(info.filename, src.read(info)))
    path.write_bytes(out.getvalue())


def test_demo_xlsx_and_csv_match():
    parent = demo_profile().parent
    a, b = load_table(parent / "clients.xlsx"), load_table(parent / "clients.csv")
    assert a.headers == b.headers and a.records == b.records
    assert a.records[0].values["Номер"] == "001"


@pytest.mark.parametrize("encoding,delimiter", [("utf-8-sig", ";"), ("cp1251", ","), ("utf-8-sig", "\t")])
def test_csv_quoted_and_multiline(tmp_path, encoding, delimiter):
    p = tmp_path / "данные.csv"
    p.write_bytes((f'Клиент{delimiter}Текст\n"Иван, Анна"{delimiter}"строка 1\nстрока 2"\n').encode(encoding))
    t = load_table(p)
    assert t.records[0].values["Текст"] == "строка 1\nстрока 2"


@pytest.mark.parametrize("content", ["A;A\n1;2", "A;\n1;2", "A;B\n1;2;3", "A;B\n", ""])
def test_invalid_csv(tmp_path, content):
    p = tmp_path / "bad.csv"
    p.write_text(content, encoding="utf-8")
    with pytest.raises(DocFlowError):
        load_table(p)


def test_short_csv_row_and_empty_row_numbers(tmp_path):
    p = tmp_path / "a.csv"
    p.write_text("A;B\n\nhello\n", encoding="utf-8")
    table = load_table(p)
    assert table.records[0].row_number == 3
    assert table.records[0].values == {"A": "hello", "B": ""}


def test_xlsx_formula_rejected(tmp_path):
    p = tmp_path / "formula.xlsx"
    wb = Workbook()
    wb.active.append(["A"])
    wb.active.append(["=1+1"])
    wb.save(p)
    with pytest.raises(DocFlowError, match="Формула"):
        load_table(p)


def test_xlsx_formatted_identifier_and_date(tmp_path):
    from datetime import date
    p = tmp_path / "a.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["id", "date"])
    ws.append([7, date(2026, 10, 21)])
    ws["A2"].number_format = "0000"
    wb.save(p)
    assert load_table(p).records[0].values == {"id": "0007", "date": "2026-10-21"}


@pytest.mark.parametrize("kind,value", [("number", "NaN"), ("number", "-1"), ("number", "Infinity"), ("number", "hello"), ("email", "a@"), ("date", "2026-02-30"), ("date", "21.10.2026")])
def test_validation_rejects_invalid(kind, value):
    _, issues = values_for(Record(2, {"col":value}), {"x":"col"}, {"x":Rule(kind=kind)})
    assert len(issues) == 1
    assert value not in issues[0].message


@pytest.mark.parametrize("kind,value", [("number", "1 250,25"), ("number", "0"), ("email", "a@example.com"), ("date", "2026-10-21")])
def test_validation_accepts_valid(kind, value):
    _, issues = values_for(Record(2, {"col":value}), {"x":"col"}, {"x":Rule(kind=kind)})
    assert not issues


def test_optional_blank():
    assert not values_for(Record(2, {"col":""}), {"x":"col"}, {"x":Rule(False,"email")})[1]
    assert values_for(Record(2, {"col":""}), {"x":"col"}, {})[1]


def test_split_runs_and_preserve_formatting(tmp_path):
    p = tmp_path / "split.docx"
    d = Document()
    para = d.add_paragraph("Before ")
    para.add_run("{{na").bold = True
    para.add_run("me}}")
    para.add_run(" / {{name}} after")
    d.save(p)
    raw = Template.load(p).render({"name":"Анна & <Иван>"})
    generated = Document(io.BytesIO(raw))
    assert generated.paragraphs[0].text == "Before Анна & <Иван> / Анна & <Иван> after"
    assert generated.paragraphs[0].runs[1].bold
    assert "{{" not in generated.paragraphs[0].text


def test_headers_footers_and_nested_table(tmp_path):
    p = tmp_path / "stories.docx"
    d = Document()
    d.add_paragraph("{{name}}")
    d.sections[0].header.paragraphs[0].text = "Header {{name}}"
    d.sections[0].footer.paragraphs[0].text = "Footer {{name}}"
    cell = d.add_table(rows=1, cols=1).cell(0,0)
    cell.add_table(rows=1, cols=1).cell(0,0).text = "{{name}}"
    d.save(p)
    t = Template.load(p)
    text = t.preview({"name":"Тест"})
    assert "Header Тест" in text and "Footer Тест" in text
    assert text.count("Тест") == 4


def test_newlines_and_tabs_in_values(tmp_path):
    t = Template.load(template(tmp_path))
    raw = t.render({"name":"one\ntwo\tthree"})
    assert Document(io.BytesIO(raw)).paragraphs[0].text == "one\ntwo\tthree"


def test_inserted_braces_are_literal(tmp_path):
    t = Template.load(template(tmp_path))
    assert "{{other}}" in t.preview({"name":"{{other}}"})


@pytest.mark.parametrize("text", ["{{user.name}}", "{{ 7 * 7 }}", "{% for x in rows %}", "{{unclosed", "no fields", "{{x|upper}}", "{# hidden #}"])
def test_template_language_not_evaluated(tmp_path, text):
    with pytest.raises(DocFlowError):
        Template.load(template(tmp_path, text))


def test_xml_entity_payload_rejected(tmp_path):
    p = template(tmp_path)
    rewrite(p, lambda name, data: b'<!DOCTYPE foo [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><foo>&xxe;</foo>' if name == "word/document.xml" else data)
    with pytest.raises(DocFlowError, match="небезопасный XML"):
        Template.load(p)


def test_external_links_rejected(tmp_path):
    p = tmp_path / "link.docx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("_rels/.rels", '<Relationships><Relationship TargetMode="External" Target="https://example.com"/></Relationships>')
    with pytest.raises(DocFlowError, match="Внешние"):
        Template.load(p)


@pytest.mark.parametrize("name", ["../word/document.xml", "word/embeddings/evil.bin", "word/vbaProject.bin", "word/activeX/a.bin", "C:\\evil.xml"])
def test_dangerous_archive_entries_rejected(name):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr(name, b"x")
    with pytest.raises(DocFlowError):
        check_archive(b.getvalue())


def test_word_fields_rejected(tmp_path):
    p = template(tmp_path)
    rewrite(p, lambda name, data: data.replace(b"</w:body>", b'<w:p><w:r><w:instrText>DDE cmd</w:instrText></w:r></w:p></w:body>') if name == "word/document.xml" else data)
    with pytest.raises(DocFlowError, match="Поля Word"):
        Template.load(p)


def test_incomplete_word_package_rejected(tmp_path):
    p = tmp_path / "incomplete.docx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>{{name}}</w:t></w:r></w:p></w:body></w:document>')
    with pytest.raises(DocFlowError):
        Template.load(p)


def test_controls_in_data_rejected(tmp_path):
    p = tmp_path / "a.csv"
    p.write_bytes(b"name\nhello\x00world")
    with pytest.raises(DocFlowError):
        load_table(p)


@pytest.mark.parametrize("name", ["../CON", "CON.txt", "LPT1", "a/b:c*?", " . ", "\u202eabc"])
def test_windows_safe_names(name):
    result = safe_name(name)
    assert result and "/" not in result and "\\" not in result and ":" not in result
    assert result.split(".")[0].upper() not in {"CON", "LPT1"}
    assert not result.endswith(("."," ")) and "\u202e" not in result


def test_reports_are_csv_injection_safe():
    assert csv_safe("=cmd") == "'=cmd"
    assert csv_safe("   +formula").startswith("'")


def test_batch_two_documents_per_valid_row_and_unchanged_sources(tmp_path):
    profile = Profile.load(demo_profile())
    before = {p:p.read_bytes() for p in [profile.data_file]+profile.templates}
    prepared = prepare(profile)
    a = generate(prepared, tmp_path)
    b = generate(prepared, tmp_path)
    assert a.directory != b.directory
    assert (a.total,a.accepted,a.rejected,a.documents) == (3,2,1,4)
    assert len(list(a.directory.rglob("*.docx"))) == 4
    assert not (a.directory / "row_0004_003").exists()
    report = json.loads((a.directory / "report.json").read_text())
    assert report["rejected"] == 1 and report["records"][2]["status"] == "error"
    assert all(p.read_bytes() == raw for p,raw in before.items())
    assert not list(tmp_path.glob(".docflow-*"))
    assert "client-a@example.com" not in (a.directory / "report.json").read_text()


def test_failed_second_template_removes_first_doc(tmp_path, monkeypatch):
    prepared = prepare(Profile.load(demo_profile()))
    original = Template.render
    def broken(self, values):
        if self.path.name.startswith("02"):
            raise ValueError("private value must not leak")
        return original(self, values)
    monkeypatch.setattr(Template, "render", broken)
    result = generate(prepared, tmp_path)
    assert result.documents == 0 and result.rejected == 3
    assert not list(result.directory.rglob("*.docx"))
    assert "private value" not in (result.directory / "report.json").read_text()


def test_disk_failure_publishes_no_partial_run(tmp_path, monkeypatch):
    import docflow.batch as batch
    prepared = prepare(Profile.load(demo_profile()))
    monkeypatch.setattr(batch.os, "fsync", lambda fd: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(DocFlowError, match="сохранить"):
        generate(prepared, tmp_path)
    assert not list(tmp_path.iterdir())


def test_profile_roundtrip(tmp_path):
    profile = Profile.load(demo_profile())
    target = tmp_path / "settings.json"
    profile.save(target)
    loaded = Profile.load(target)
    assert loaded.data_file == profile.data_file
    assert loaded.mapping == profile.mapping and loaded.rules == profile.rules
    prepare(loaded)


def test_profile_cannot_overwrite_input():
    profile = Profile.load(demo_profile())
    before = profile.data_file.read_bytes()
    with pytest.raises(DocFlowError):
        profile.save(profile.data_file)
    assert profile.data_file.read_bytes() == before


def test_missing_mapping_is_actionable():
    profile = Profile.load(demo_profile())
    del profile.mapping["client"]
    with pytest.raises(DocFlowError, match="client"):
        prepare(profile)


def test_cli_demo_and_preview(tmp_path, capsys):
    assert main(["demo", "--out", str(tmp_path)]) == 2
    assert json.loads(capsys.readouterr().out)["documents"] == 4
    assert main(["preview",str(demo_profile())]) == 0
    assert "Демо Клиент А" in capsys.readouterr().out

