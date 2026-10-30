"""Behavior regressions for the real Qt workspace, not screenshots/mock counters."""
import csv
import io
import threading
import time
import zipfile
from pathlib import Path

import pytest
from docx import Document
from PySide6.QtCore import QTimer, Qt, QPoint
from PySide6.QtGui import QCloseEvent
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from docflow.cli import demo_profile
from docflow.data import load_table
from docflow.errors import DocFlowError
from docflow.gui import MainWindow
from docflow.profile import Profile
from docflow.release_check import wait_idle
from docflow.safety import check_archive
from docflow.templates import Template
from docflow.ui import window as workspace


@pytest.fixture
def window():
    app=QApplication.instance() or QApplication([])
    widget=MainWindow();widget.show();widget.load_demo();wait_idle(widget)
    yield widget
    wait_idle(widget);widget.close();widget.deleteLater();app.processEvents()


def pump_until(predicate,timeout=5):
    deadline=time.monotonic()+timeout
    while not predicate():
        QApplication.processEvents();time.sleep(.005)
        if time.monotonic()>deadline:raise AssertionError('Timed out')


def test_navigation_preserves_loaded_state_and_field_rules(window):
    original=window.table;mapping,rules=window.mapping()
    window.mapping_table.cellWidget(0,3).setChecked(False)
    for page in [1,2,3,4,0]:window.navigate(page)
    assert window.table is original and window.mapping()[0]==mapping
    assert window.mapping_table.cellWidget(0,3).isChecked() is False
    assert window.data_model.rowCount()==3 and len(window.templates)==2


def test_wide_navigation_labels_fit_without_truncation(window):
    window.resize(1920,1080);QApplication.processEvents()
    for button in window.nav_buttons:
        assert button.width()>=button.sizeHint().width(),button.text()


def test_template_selection_preserves_mapping_rules_and_folder_column(window,tmp_path):
    before=window.mapping();folder=window.name_selector.currentText()
    window.template_list.item(0).setCheckState(Qt.CheckState.Unchecked)
    assert len(window.active_templates())==1
    window.template_list.item(0).setCheckState(Qt.CheckState.Checked)
    wait_idle(window)
    assert window.mapping()==before and window.name_selector.currentText()==folder
    window.output_path.setText(str(tmp_path));window.start_generation();wait_idle(window)
    assert window.last_result.documents==4


def test_blank_mapping_not_silently_restored_and_stale_preview_clears(window):
    window.mapping_table.cellWidget(0,1).setCurrentText('')
    window.template_list.item(0).setCheckState(Qt.CheckState.Unchecked)
    window.template_list.item(0).setCheckState(Qt.CheckState.Checked)
    assert window.mapping_table.cellWidget(0,1).currentText()==''
    assert not window.generate_btn.isEnabled()
    assert 'Не сопоставлены' in window.preview_text.toPlainText()
    while window.active_templates():
        for i,t in enumerate(window.templates):
            if t.path in window._active_paths:
                window.template_list.item(i).setCheckState(Qt.CheckState.Unchecked);break
    assert 'Загрузите таблицу' in window.preview_text.toPlainText()


def test_selected_templates_only_are_generated(window,tmp_path):
    window.template_list.item(1).setCheckState(Qt.CheckState.Unchecked)
    window.output_path.setText(str(tmp_path));window.start_generation();wait_idle(window)
    assert window.last_result.documents==2 and window.last_result.accepted==2


def test_large_data_model_exposes_all_rows_and_plain_full_tooltips(window,tmp_path):
    base=Profile.load(demo_profile());source=tmp_path/'5000 записей.csv'
    with source.open('w',encoding='utf-8',newline='') as f:
        writer=csv.writer(f);writer.writerow(['Клиент','Номер','Дата','Услуга','Сумма','Email'])
        writer.writerows([[f'Клиент {i}','001','2026-10-21','Сервис','100',f'{i}@example.com'] for i in range(5000)])
    window.apply_profile(Profile(source,base.templates,base.mapping,base.rules,base.name_column));wait_idle(window)
    assert window.source_table.rowCount()==5000 and window.row_selector.count()==5000
    model=window.source_table.model()
    assert model.data(model.index(4999,0))=='Клиент 4999'
    window.row_selector.setCurrentIndex(4999);wait_idle(window)
    assert 'Клиент 4999' in window.preview_text.toPlainText()


def test_file_loading_runs_in_worker_and_failure_keeps_existing_state(window,tmp_path,monkeypatch):
    entered=threading.Event();release=threading.Event();old=window.table
    def blocked(*args,**kwargs):
        entered.set()
        if not release.wait(3):raise RuntimeError('test gate timeout')
        raise DocFlowError('Повреждённая таблица')
    monkeypatch.setattr(workspace,'load_table',blocked)
    monkeypatch.setattr(QFileDialog,'getOpenFileName',lambda *a:('synthetic.csv',''))
    ticks=[];timer=QTimer();timer.setInterval(5);timer.timeout.connect(lambda:ticks.append(True));timer.start()
    window.select_data();pump_until(entered.is_set);pump_until(lambda:len(ticks)>=4)
    assert window._busy and window.table is old
    release.set();wait_idle(window);timer.stop()
    assert window.table is old and not window._busy
    assert 'Повреждённая' in window.status.text()


def test_stale_worker_preview_cannot_replace_new_selection(window,monkeypatch):
    entered=threading.Event();release=threading.Event();original=Template.preview
    def blocked(template,values):
        if values['client']=='Демо Клиент А':
            entered.set()
            if not release.wait(3):raise RuntimeError('test gate timeout')
        return original(template,values)
    monkeypatch.setattr(Template,'preview',blocked)
    window.row_selector.setCurrentIndex(0);window.update_preview();pump_until(entered.is_set)
    window.row_selector.setCurrentIndex(1);release.set();wait_idle(window)
    assert 'Демо «Вектор»' in window.preview_text.toPlainText()
    assert 'Демо Клиент А' not in window.preview_text.toPlainText()


def test_double_generation_and_close_during_work(window,tmp_path,monkeypatch):
    original=workspace.generate;entered=threading.Event();release=threading.Event();calls=[]
    def blocked(*args,**kwargs):
        calls.append(1);entered.set()
        if not release.wait(3):raise RuntimeError('test gate timeout')
        return original(*args,**kwargs)
    monkeypatch.setattr(workspace,'generate',blocked)
    monkeypatch.setattr(QMessageBox,'information',lambda *args:QMessageBox.StandardButton.Ok)
    window.output_path.setText(str(tmp_path));window.start_generation();first=window.worker
    pump_until(entered.is_set);window.start_generation()
    assert window.worker is first and len(calls)==1
    event=QCloseEvent();window.closeEvent(event);assert not event.isAccepted()
    release.set();wait_idle(window)
    assert window.last_result.documents==4 and window.generate_btn.isEnabled()


def test_worker_error_unlocks_controls_and_publishes_no_success(window,tmp_path,monkeypatch):
    def fail(*args,**kwargs):raise DocFlowError('Диск недоступен')
    monkeypatch.setattr(workspace,'generate',fail)
    window.output_path.setText(str(tmp_path));window.start_generation();wait_idle(window)
    assert not window._busy and window.generate_btn.isEnabled()
    assert window.last_result is None and 'Диск недоступен' in window.status.text()
    assert not list(tmp_path.iterdir())


def test_profile_roundtrip_after_ui_changes(window,tmp_path,monkeypatch):
    window.mapping_table.cellWidget(0,3).setChecked(False)
    expected=window.mapping();profile=tmp_path/'Настройки с пробелами.json'
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a:(str(profile),''))
    window.save_profile();assert profile.exists()
    monkeypatch.setattr(QFileDialog,'getOpenFileName',lambda *a:(str(profile),''))
    window.open_profile();wait_idle(window)
    assert window.mapping()==expected


def test_snapshot_copies_mutable_data_but_keeps_immutable_template_bytes(window):
    prepared=window.snapshot()
    original=prepared.table.records[0].values['Клиент']
    window.table.records[0].values['Клиент']='Изменено после снимка'
    assert prepared.table.records[0].values['Клиент']==original
    assert prepared.table is not window.table
    assert prepared.templates[0].raw is window.active_templates()[0].raw


@pytest.mark.parametrize('size',[(640,360),(1024,750),(1280,720),(1400,860)])
def test_primary_action_remains_reachable_when_window_resizes(window,size):
    window.resize(*size);QApplication.processEvents()
    for page in range(4):
        window.navigate(page);QApplication.processEvents()
        primary=window.generate_btn if page==3 else window.core_next;point=primary.mapTo(window,QPoint())
        assert primary.isVisible() and primary.isEnabled() and window.rect().contains(point)
        assert window.rect().contains(point+QPoint(primary.width()-1,primary.height()-1))


@pytest.mark.parametrize('value',['literal {{not_a_field}} \\ text','line 1\r\nline 2\tvalue','<tag> & "quotes"'])
def test_fast_preview_preserves_literal_values_and_rendered_content(tmp_path,value):
    path=tmp_path/'split.docx';d=Document();p=d.add_paragraph();p.add_run('Before {{na').bold=True;p.add_run('me}} after');d.save(path)
    template=Template.load(path)
    preview=template.preview({'name':value})
    rendered=Document(io.BytesIO(template.render({'name':value}))).paragraphs[0].text
    assert preview==rendered
    assert Document(io.BytesIO(template.render({'name':value}))).paragraphs[0].runs[0].bold


@pytest.mark.parametrize('value',[123,'\x00bad','\ud800bad'])
def test_preview_retains_input_safety_checks(tmp_path,value):
    path=tmp_path/'safe.docx';d=Document();d.add_paragraph('{{name}}');d.save(path)
    with pytest.raises(DocFlowError):Template.load(path).preview({'name':value})


@pytest.mark.parametrize('rows,columns',[(5001,1),(1,101)])
def test_table_limits_remain_enforced(tmp_path,rows,columns):
    path=tmp_path/'limit.csv'
    with path.open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow([f'column{i}' for i in range(columns)])
        writer.writerows([['value']*columns for _ in range(rows)])
    with pytest.raises(DocFlowError):load_table(path)


def test_archive_member_limit_not_weakened():
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w') as z:
        for i in range(2001):z.writestr(f'{i}.bin',b'')
    with pytest.raises(DocFlowError):check_archive(b.getvalue())


@pytest.mark.parametrize('size',[(1024,576),(900,600),(640,360)])
def test_create_page_does_not_gain_horizontal_scroll_at_compact_width(window,size):
    window.resize(*size);window.navigate(0);QApplication.processEvents();QApplication.processEvents()
    assert window.pages.currentWidget().horizontalScrollBar().maximum()==0
